"""장 순서대로 ReportSheet 에 써 넣는다. 두 패스(1패스로 장·절 쪽번호 확정 → 2패스로 목차 채움)."""
from __future__ import annotations

import os
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .collector import ReportData, _fmt
from .sheet import FONT, LIGHT, ReportSheet

LOGO_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "templates", "hd_logo.png"))
STATUS_KO = {"pass": "PASS", "warn": "WARN", "fail": "FAIL", "skip": "SKIP"}


def toc_entries(d: ReportData) -> list[tuple[str, int]]:
    """(제목, level). level 0 = 장. 절 제목은 _write 와 글자 단위로 같아야 쪽번호가 채워진다."""
    e = [("요약 (Summary)", 0),
         ("1. 개요", 0), ("1.1 목적", 1), ("1.2 검토 범위", 1), ("1.3 입력 자료", 1), ("1.4 소프트웨어", 1), ("1.5 판정 기준", 1),
         ("2. 해석 모델", 0), ("2.1 모델 개요", 1), ("2.2 중량 및 무게중심", 1), ("2.3 재료 및 단면", 1), ("2.4 모델 편집 이력", 1),
         ("3. 권상 위치 선정 및 자세 안정성", 0), ("3.1 권상 그룹 및 러그 절점", 1), ("3.2 권상 위치 선정 근거", 1),
         ("3.3 자세 안정성 단계별 판정", 1), ("3.4 정점 및 와이어 기하", 1), ("3.5 전도 여유", 1),
         ("4. 하중 및 경계조건", 0), ("4.1 하중", 1), ("4.2 와이어 모델링", 1), ("4.3 경계조건", 1),
         ("4.4 안정화 파라미터", 1), ("4.5 해석 조건 및 진단", 1),
         ("5. 해석 결과", 0), ("5.1 변위", 1), ("5.2 부재 응력", 1), ("5.3 와이어 장력 및 반력", 1)]
    if d.edits and d.edits.support_beams:
        e.append(("5.4 가서포트", 1))
    e += [("6. 결론", 0), ("부록 A. 단면 물성", 0), ("부록 B. 기호", 0), ("부록 C. 해석 로그 요약", 0)]
    return e


def build_workbook(d: ReportData, figs: dict[str, bytes], fig_errors: Optional[dict[str, str]] = None) -> Workbook:
    fig_errors = fig_errors or {}
    logo = None
    if os.path.isfile(LOGO_PATH):
        with open(LOGO_PATH, "rb") as fh:
            logo = fh.read()
    entries = toc_entries(d)
    # 1패스: 장·절 쪽번호와 총 페이지 수를 센다. 2패스에서 목차와 바닥글(Page n / N)을 채운다.
    pages, total = _write(Workbook(), d, figs, fig_errors, logo, entries, None, None)
    wb = Workbook()
    _write(wb, d, figs, fig_errors, logo, entries, pages, total)
    _data_sheets(wb, d)
    return wb


def _write(wb, d: ReportData, figs, errs, logo, entries, pages, total_pages):
    i, r, v, st, m, L = d.identity, d.results, d.verdicts, d.stability, d.model, d.loads
    s = ReportSheet(
        wb.active,
        doc_title=f"{i.title_prefix} 권상 구조 검토 보고서",
        hull=i.hull_no, unit=i.unit_no, method=i.lifting_method,
        department=i.department if i.department != "-" else "", drawing_no=i.drawing_no,
        logo_png=logo, total_pages=total_pages,
    )
    sec_pages: dict[str, int] = {}

    def sec(text):
        sec_pages[text] = s.section(text)

    def chap(key, text):
        sec_pages[text] = s.chapter(key, text)

    def fig(key, caption, w=0, h=0):
        if key in figs:
            s.figure(figs[key], caption, w, h)
        elif key in errs:
            s.placeholder(caption, errs[key])

    def none_note(what):
        s.note(f"자료 없음 — {what} 파일이 결과 폴더에 없어 이 절은 생략되었습니다.")

    # ── 표지 ────────────────────────────────────────────────────────────────
    s.cover(f"{i.title_prefix} 권상 구조 검토 보고서",
            [("HULL NO.", i.hull_no), ("UNIT NO.", i.unit_no), ("도면 번호", i.drawing_no),
             ("권상 방식", f"{i.lifting_method} ({i.equipment or 'Hook'})"), ("리비전", f"Rev. {i.revision}"),
             ("작성", f"{i.author} / {i.department}"), ("작성일", i.generated_at)],
            [("구조 판정", v.structure, "ok" if v.structure == "OK" else ("ng" if v.structure == "NG" else "skip")),
             ("자세 안정성", STATUS_KO.get(v.stability, str(v.stability).upper()), v.stability),
             ("국부변형 방지 지그", "필요" if v.jig_required else "불요", "warn" if v.jig_required else "ok")])

    # ── 요약 ────────────────────────────────────────────────────────────────
    chap("summary", "요약 (Summary)")
    s.table(["중량 [ton]", "항복강도 [MPa]", "허용응력 [MPa]", "최대응력 [MPa]", "최대변위 [mm]", "판정"],
            [[round(m.total_mass_ton, 3), round(r.yield_mpa, 1), round(r.allowable_mpa, 1),
              f"{r.max_stress_mpa:.2f} (E{r.max_stress_element_id})",
              f"{r.max_displacement_mm:.2f} (N{r.max_displacement_node_id})", v.structure]],
            widths=[2, 2, 2, 2, 2, 2], align=["center"] * 6, status_col=5, title="검토 요약")
    s.note(f"허용응력 = 항복강도 × 0.8. 하중 = 자중 × 안전계수 {L.safety_factor:g}. 변위는 참고치(판정 없음).")
    s.table(["Hook/Trolley", "와이어 수", "장력 합 [ton]", "수직 반력 합 [ton]"],
            [[f"G{h.group_id}", h.wire_count, round(h.tension_ton, 3), round(h.vertical_ton, 3)] for h in r.hook_totals]
            + [["합계", sum(h.wire_count for h in r.hook_totals), round(sum(h.tension_ton for h in r.hook_totals), 3),
                round(sum(h.vertical_ton for h in r.hook_totals), 3)]],
            widths=[3, 3, 3, 3], title="Hook 별 반력 요약")
    if r.hook_check_error_pct is not None:
        s.note(f"검산: Σ수직반력 {sum(h.vertical_ton for h in r.hook_totals):.3f} ton vs 설계하중 {L.total_load_ton:.3f} ton "
               f"(오차 {r.hook_check_error_pct:+.1f} %).")
    s.subsection("먼저 확인할 사항")
    for a in (v.actions[:3] or ["조치가 필요한 항목이 없습니다."]):
        s.bullet(a)

    # ── 목차 ────────────────────────────────────────────────────────────────
    s.toc_placeholder(len(entries))

    # ── 1. 개요 ─────────────────────────────────────────────────────────────
    chap("1", "1. 개요")
    sec("1.1 목적")
    s.para(f"본 보고서는 {i.title_prefix}(HULL {i.hull_no}, UNIT {i.unit_no})의 {i.lifting_method} 권상 시 구조 안전성을 "
           "검토한 결과를 기술한다. 권상 위치 선정과 자세 안정성 평가, 와이어를 포함한 선형 정적 구조 해석의 입력·가정·결과를 "
           "한 문서로 정리하여 검토자가 별도 자료 없이 판정 근거를 확인할 수 있게 한다.")
    sec("1.2 검토 범위")
    s.kv([("권상 방식", f"{i.lifting_method} ({i.equipment or 'Hook'})"), ("권상 그룹 수", f"{i.group_count} (Hook/Trolley)"),
          ("와이어 길이", f"{i.wire_length_m:g} m"), ("안전계수", f"{L.safety_factor:g}"), ("해석 종류", L.solver)])
    sec("1.3 입력 자료")
    s.kv([("원본 BDF", i.source_bdf or "-"), ("해석 모델(편집본)", i.edited_model or "-"),
          ("결과 생성 시각", i.generated_at), ("비고", i.notes or "-")])
    sec("1.4 소프트웨어")
    s.kv([("플랫폼", "Hi-TESS WorkBench / Module Unit Studio"), ("자세안정성 엔진", i.engine_version or "ModuleAnalysis.Stability"),
          ("BDF 생성·결과 매핑", "NastranBridge (LiftingBdfBuild / LiftingResultMapper)"), ("해석기", L.solver)])
    sec("1.5 판정 기준")
    s.table(["항목", "기준", "비고"],
            [["부재 응력", f"σmax ≤ σ허용 = σy × 0.8 = {r.allowable_mpa:g} MPa", f"σy = {r.yield_mpa:g} MPa"],
             ["와이어 장력", f"장력 > {d.options.jig_limit_ton:g} ton 이면 국부 변형 방지 지그 필요", "장력 = |축력| / 9,800 N/ton"],
             ["와이어 건전성", "압축·슬랙 와이어 없어야 함", "CROD 축력 부호"],
             ["자세 안정성", "엔진 7단계 종합 판정 (FAIL 이면 권상 불가)", "Strict 평가 " + ("ON" if st.strict_evaluation else "OFF")],
             ["변위", "참고치 (판정 없음)", "최대 변위 절점·성분 기록"]], widths=[3, 6, 3], title="판정 기준")

    # ── 2. 해석 모델 ─────────────────────────────────────────────────────────
    chap("2", "2. 해석 모델")
    sec("2.1 모델 개요")
    s.kv([("절점 / 요소", f"{m.node_count:,} / {m.element_count:,} (RBE2 {m.rigid_count:,}, 집중질량 {m.point_mass_count:,})"),
          ("단면(PID) / 재료", f"{m.property_count} / {m.material_count}"),
          ("외형 치수 (X × Y × Z)", f"{m.size_mm[0]:,.0f} × {m.size_mm[1]:,.0f} × {m.size_mm[2]:,.0f} mm"),
          ("좌표 범위", f"X {m.bbox['minX']:,.0f}~{m.bbox['maxX']:,.0f} · Y {m.bbox['minY']:,.0f}~{m.bbox['maxY']:,.0f}"
                        f" · Z {m.bbox['minZ']:,.0f}~{m.bbox['maxZ']:,.0f}"),
          ("입력 검증", f"{(m.validation_status or '-').upper()} · 자유단 절점 "
                       f"{m.free_end_nodes if m.free_end_nodes is not None else '-'} · 분리 그룹 "
                       f"{m.disconnected_groups if m.disconnected_groups is not None else '-'}")])
    fig("model_plan", "해석 모델 평면도 (XY)")
    fig("model_side", "해석 모델 측면도 (XZ)")
    fig("model_iso", "해석 모델 등각도")
    sec("2.2 중량 및 무게중심")
    c = m.cog_mm
    s.kv([("총 중량", f"{m.total_mass_ton:.3f} ton ({m.mass_source or '-'})"),
          ("무게중심 (X, Y, Z)", f"({c.get('x', 0):,.1f}, {c.get('y', 0):,.1f}, {c.get('z', 0):,.1f}) mm")])
    if m.point_masses:
        s.table(["CONM2 ID", "절점", "질량 [ton]"],
                [[p.get("id"), p.get("nodeId"), float(p.get("mass") or 0)] for p in m.point_masses[:40]],
                widths=[4, 4, 4], title=f"집중질량 목록 (상위 {min(40, len(m.point_masses))}개 / 총 {len(m.point_masses)})")
    sec("2.3 재료 및 단면")
    s.table(["MID", "이름", "E [MPa]", "ν", "ρ [ton/mm³]"],
            [[x.mid, x.name, x.e_mpa, x.nu, x.rho] for x in m.materials], widths=[2, 3, 3, 2, 2], title="재료")
    s.table(["PID", "단면", "치수 [mm]", "면적 [mm²]", "MID"],
            [[x.pid, f"{x.card} {x.kind}", ", ".join(_fmt(vv) for vv in x.dims),
              round(x.area_mm2, 1) if x.area_mm2 else "-", x.material_id] for x in m.sections],
            widths=[1, 3, 4, 2, 2], title="단면 목록")
    sec("2.4 모델 편집 이력")
    if d.edits:
        e = d.edits
        s.kv([("삭제 요소 / 강체 / 절점", f"{e.deleted_element_count} / {e.deleted_rigid_count} / {e.deleted_node_count}"),
              ("가서포트 추가", f"{len(e.support_beams)}개"),
              ("제거된 원본 경계조건", f"SPC 계열 {e.removed_spc_cards}장 · SUPORT {e.removed_suport_cards}장 "
                                  "(권상 해석은 정점 SPC 만 사용)")])
        if e.support_beams:
            s.table(["No.", "요소 ID", "절점 A", "절점 B", "길이 [mm]", "단면 치수"],
                    [[k + 1, sb.element_id, sb.start_node, sb.end_node, round(sb.length_mm, 0),
                      ", ".join(_fmt(vv) for vv in sb.dims)] for k, sb in enumerate(e.support_beams)],
                    widths=[1, 2, 2, 2, 2, 3], title="가서포트 목록")
    else:
        none_note("편집 모델(_edited.json) 또는 원본 json")

    # ── 3. 권상 위치·자세 안정성 ─────────────────────────────────────────────
    h = d.hoist
    chap("3", "3. 권상 위치 선정 및 자세 안정성")
    sec("3.1 권상 그룹 및 러그 절점")
    if h:
        rows = []
        for g in h.groups:
            for n in (g.nodes or [{"id": nid} for nid in g.node_ids]):
                rows.append([f"G{g.group_id}", n.get("id"), n.get("x", "-"), n.get("y", "-"), n.get("z", "-")])
        s.table(["그룹", "러그 절점", "X [mm]", "Y [mm]", "Z [mm]"], rows, widths=[2, 2, 3, 3, 2], title="권상 그룹별 러그 절점")
    else:
        s.table(["그룹", "러그 절점"],
                [[f"G{a.get('groupId')}", ", ".join(str(n) for n in a.get("assignedNodeIds", []))] for a in st.apexes],
                widths=[3, 9], title="권상 그룹")
    fig("hoist_plan", "권상 배치 평면도 — 러그·정점·와이어·COG")
    fig("hoist_side", "권상 배치 측면도")
    fig("hoist_iso", "권상 위치 3D — 지시 핀은 권상 그룹, 검정은 무게중심")
    sec("3.2 권상 위치 선정 근거")
    if h:
        if h.auto_selected:
            s.para(f"권상 위치는 엔진 최적화(후보 {h.candidate_count:,}개 조립, {h.evaluated_count:,}개 평가, "
                   f"{h.elapsed_ms / 1000:.1f} s)로 자동 선정되었다. 평가 결과 PASS {h.pass_count:,} / WARN {h.warn_count:,} / "
                   f"FAIL {h.fail_count:,} 중 '{h.best_label}' (점수 {h.best_score:,.0f}) 이 선택되었다.")
        else:
            s.para("권상 위치는 사용자가 Studio 에서 직접 지정하였다(최적화 결과와 다름). 아래 표는 참고용 최적화 지표다.")
        keys = [("stage6Status", "전도 판정"), ("stage6MarginMm", "전도 여유 [mm]"), ("minSlingAngleDeg", "최소 슬링각 [°]"),
                ("wireConflictCount", "와이어 간섭 [건]"), ("supportSpanFraction", "지지 폭 비율"), ("evaluationMode", "평가 방식")]
        s.table(["지표", "최적(선정)", "이전(current)"],
                [[label, h.best_metrics.get(k, "-"), h.current_metrics.get(k, "-")] for k, label in keys],
                widths=[4, 4, 4], title="선정 지표 비교")
    else:
        none_note("권상 위치 최적화(_hoist_optimization.json)")
    sec("3.3 자세 안정성 단계별 판정")
    s.para(f"종합 판정: {STATUS_KO.get(st.overall_status, st.overall_status)} — {st.primary_message}")
    if st.shape_gate_relaxed:
        s.note("Strict 평가 OFF: 1·2단계 형상 판정이 FAIL 대신 WARN 으로 완화되어 평가되었다"
               "(3단계 wire≤0, 6단계 전도는 완화 없음).")
    s.table(["단계", "항목", "판정", "핵심 수치"],
            [[x.stage, x.label, STATUS_KO.get(x.status, x.status),
              x.key_metric + (f"  —  {x.note}" if x.note else "")] for x in st.stage_rows],
            widths=[1, 3, 1, 7], status_col=2, title="자세 안정성 7단계 판정")
    for w in st.warnings:
        s.bullet(f"[{w.get('stage')}단계 · {w.get('target', '')}] {w.get('headline', '')} — {w.get('message', '')}")
    sec("3.4 정점 및 와이어 기하")
    s.table(["그룹", "정점 X", "정점 Y", "정점 Z", "연결 러그"],
            [[f"G{a.get('groupId')}", a["pointMm"]["x"], a["pointMm"]["y"], a["pointMm"]["z"],
              ", ".join(f"N{n}" for n in a.get("assignedNodeIds", []))] for a in st.apexes],
            widths=[1, 2, 2, 2, 5], title="권상 정점(Hook/Trolley) 좌표 [mm]")
    s.table(["그룹", "러그", "길이 [mm]", "슬링각 [°]", "상태"],
            [[f"G{w.group_id}", f"N{w.lug_node_id}", round(w.length_mm, 1), round(w.angle_deg, 2),
              "pass" if w.safe else "warn"] for w in st.wires],
            widths=[2, 2, 3, 3, 2], status_col=4,
            title=f"와이어 기하 (최소 슬링각 {st.min_sling_angle_deg if st.min_sling_angle_deg is not None else '-'}°)")
    sec("3.5 전도 여유")
    if st.tipping:
        t = st.tipping
        s.kv([("평가 방식", t.evaluation_mode), ("안정 여부", "안정" if t.is_stable else "불안정"),
              ("전도 여유 / 내부 여유", f"{t.margin_mm:,.1f} / {t.interior_margin_mm:,.1f} mm"),
              ("COG 편차 / 임계", f"{t.deviation_mm:,.1f} / {t.threshold_mm:,.1f} mm"),
              ("COG 지지영역 내부", "예" if t.cog_inside else "아니오"), ("경사각", f"{t.tilt_deg:.2f}°"),
              ("정점–COG 높이", f"{t.apex_to_cog_mm:,.1f} mm")])
    else:
        none_note("6단계 전도 평가 결과")

    # ── 4. 하중·경계조건 ─────────────────────────────────────────────────────
    chap("4", "4. 하중 및 경계조건")
    sec("4.1 하중")
    s.kv([("중력 가속도", f"{L.gravity_mm_s2:g} mm/s² (−Z), GRAV SID {L.grav_sid}"),
          ("안전계수", f"{L.safety_factor:g} (LOAD SID {L.load_sid} 로 배율)"),
          ("설계 하중", f"{m.total_mass_ton:.3f} ton × {L.safety_factor:g} = {L.total_load_ton:.3f} ton")])
    sec("4.2 와이어 모델링")
    s.kv([("요소", f"CROD, PID {L.wire_pid}{' (충돌 회피 재배정)' if L.wire_pid_remapped else ''}, "
                  f"MID {L.wire_mid}{' (충돌 회피 재배정)' if L.wire_mid_remapped else ''}"),
          ("단면", f"A = {L.wire_area_mm2:,.0f} mm², J = {L.wire_j_mm4:,.0f} mm⁴"),
          ("재료", f"E = {L.wire_e_mpa:,.0f} MPa"), ("연결", "각 러그 절점 ↔ 그룹 정점 절점, 축력만 전달")])
    sec("4.3 경계조건")
    s.kv([("정점 구속", f"SPC SID {L.spc_sid}, 성분 {L.spc_components}"),
          ("강체 운동 방지 앵커",
           f"절점 {L.anchor_node_id} 성분 {L.anchor_components} (COG 로부터 {_fmt(L.anchor_distance_mm)} mm)"
           if L.anchor_node_id else "미적용"),
          ("원본 경계조건", "권상 해석에서는 원본 SPC/SUPORT 를 모두 제거하고 위 조건만 사용")])
    sec("4.4 안정화 파라미터")
    s.para(L.stabilization_reason or "-")
    for card in L.stabilization_cards:
        s.bullet(card)
    sec("4.5 해석 조건 및 진단")
    f = L.f06
    s.kv([("해석", f"{L.solver}, SUBCASE {L.subcase_id}"),
          ("F06 진단", f"FATAL {f.get('fatalCount', 0)} · ERROR {f.get('errorCount', 0)}" if f.get("available") else "F06 파일 없음")])
    for smp in (f.get("fatalSamples") or [])[:3]:
        s.note(smp)

    # ── 5. 해석 결과 ─────────────────────────────────────────────────────────
    chap("5", "5. 해석 결과")
    sec("5.1 변위")
    t1, t2, t3 = r.max_displacement_components
    s.kv([("최대 변위", f"{r.max_displacement_mm:.3f} mm @ 절점 {r.max_displacement_node_id}"),
          ("성분 (T1, T2, T3)", f"({t1:.3f}, {t2:.3f}, {t3:.3f}) mm"), ("평가", "참고치 — 판정 기준 없음")])
    s.table(["절점", "T1 [mm]", "T2 [mm]", "T3 [mm]", "크기 [mm]"],
            [[x.node_id, round(x.t1, 3), round(x.t2, 3), round(x.t3, 3), round(x.magnitude, 3)] for x in r.top_displacements],
            widths=[2, 2, 2, 2, 4], title="변위 상위 10 절점")
    fig("deformed_iso", "변형 형상 (배율 자동)")
    sec("5.2 부재 응력")
    if r.exceeding:
        s.para(f"허용응력 {r.allowable_mpa:g} MPa 를 초과하는 부재가 {len(r.exceeding)}개 있다.")
        s.table(["요소", "PID", "단면", "σmax [MPa]", "활용도", "판정"],
                [[x.element_id, x.pid, x.section_label, round(x.stress_mpa, 2), round(x.utilization or 0, 3), "NG"]
                 for x in r.exceeding], widths=[1, 1, 5, 2, 2, 1], status_col=5, title="허용응력 초과 부재 (전량)")
    else:
        s.para(f"허용응력을 초과하는 부재는 없습니다 (검사 부재 {r.member_count:,}개, "
               f"최대 활용도 {(r.governing[0].utilization if r.governing else 0):.3f}).")
    s.table(["순위", "요소", "PID", "단면", "σmax [MPa]", "활용도"],
            [[k + 1, x.element_id, x.pid, x.section_label, round(x.stress_mpa, 2), round(x.utilization or 0, 3)]
             for k, x in enumerate(r.governing)], widths=[1, 1, 1, 5, 2, 2], title=f"지배 부재 상위 {len(r.governing)}")
    s.table(["PID", "단면", "부재 수", "최대 σ [MPa]", "최대 활용도", "지배 요소"],
            [[p.pid, p.section_label, p.count, round(p.max_stress_mpa, 2), round(p.max_utilization, 3), p.governing_element_id]
             for p in r.pid_summary], widths=[1, 5, 1, 2, 2, 1], title="단면(PID)별 요약")
    fig("utilization_hist", "부재 활용도 분포", 640, 285)
    fig("stress_plan", "부재 응력 활용도 — 평면도")
    fig("stress_iso", "부재 응력 활용도 — 등각도")
    sec("5.3 와이어 장력 및 반력")
    s.table(["그룹", "러그", "축력 [N]", "장력 [ton]", "슬링각 [°]", "수직반력 [ton]", "지그", "상태"],
            [[f"G{w.group_id}", f"N{w.lug_node_id}", round(w.axial_force_n, 1), round(w.tension_ton, 3),
              round(w.angle_deg, 2) if w.angle_deg is not None else "-",
              round(w.vertical_ton, 3) if w.vertical_ton is not None else "-",
              "필요" if w.needs_jig else "불요", "NG" if (w.is_compression or w.is_slack) else "OK"] for w in r.wires],
            widths=[1, 1, 2, 2, 1, 2, 1, 2], status_col=7,
            title=f"와이어별 장력·반력 (지그 기준 {d.options.jig_limit_ton:g} ton)")
    s.table(["Hook/Trolley", "와이어 수", "장력 합 [ton]", "수직 반력 합 [ton]"],
            [[f"G{x.group_id}", x.wire_count, round(x.tension_ton, 3), round(x.vertical_ton, 3)] for x in r.hook_totals],
            widths=[3, 3, 3, 3], title="Hook 별 합계")
    if r.hook_check_error_pct is not None:
        s.note(f"검산: Σ수직반력 / (중량 × SF) − 1 = {r.hook_check_error_pct:+.1f} % "
               "(슬링각 산정과 와이어 자중 무시에 따른 차이).")
    if d.edits and d.edits.support_beams:
        sec("5.4 가서포트")
        sup_ids = {sb.element_id for sb in d.edits.support_beams}
        s.table(["요소", "단면", "σmax [MPa]", "활용도"],
                [[x.element_id, x.section_label, round(x.stress_mpa, 2), round(x.utilization or 0, 3)]
                 for x in r.all_members if x.element_id in sup_ids], widths=[2, 6, 2, 2], title="가서포트 부재 응력")
        fig("support_plan", "가서포트 배치 — 평면도")
        fig("support_iso", "가서포트 배치 — 등각도")

    # ── 6. 결론 ──────────────────────────────────────────────────────────────
    chap("6", "6. 결론")
    s.para(f"구조 판정: {v.structure} — 최대 응력 {r.max_stress_mpa:.2f} MPa (요소 {r.max_stress_element_id}) / "
           f"허용 {r.allowable_mpa:g} MPa, 초과 부재 {r.exceed_count}개.")
    s.para(f"자세 안정성: {STATUS_KO.get(v.stability, v.stability)} — {st.primary_message}")
    s.para(f"국부 변형 방지 지그: {'필요' if v.jig_required else '불요'} (기준 {d.options.jig_limit_ton:g} ton, "
           f"최대 와이어 장력 {max((w.tension_ton for w in r.wires), default=0):.3f} ton).")
    s.subsection("조치 필요 사항")
    for a in (v.actions or ["없음"]):
        s.bullet(a)
    s.subsection("근거와 한계")
    for line in ("선형 정적 해석(SOL 101)이며 동적 계수·충격은 안전계수에 포함된 것으로 간주한다.",
                 "와이어는 축력만 전달하는 CROD 로 이상화하였고 와이어 자중·탄성 신장에 따른 각도 변화는 무시하였다.",
                 "허용응력은 항복강도 × 0.8 의 사내 기준을 적용하였다. 좌굴·용접부·러그 국부 강도는 본 검토 범위 밖이다.",
                 "자세 안정성은 강체 가정의 기하 평가이며 구조 해석 결과와 독립적으로 판정한다."):
        s.bullet(line)

    # ── 부록 ─────────────────────────────────────────────────────────────────
    chap("A", "부록 A. 단면 물성")
    s.table(["PID", "카드", "종류", "치수 [mm]", "면적 [mm²]", "MID"],
            [[x.pid, x.card, x.kind, ", ".join(_fmt(vv) for vv in x.dims),
              round(x.area_mm2, 1) if x.area_mm2 else "-", x.material_id] for x in m.sections],
            widths=[1, 2, 1, 4, 2, 2])
    chap("B", "부록 B. 기호")
    s.table(["기호", "의미"],
            [["σmax", "요소 내 모든 응력점·양단 중 |σ| 최대"], ["σy / σ허용", "항복강도 / 허용응력 (= σy × 0.8)"],
             ["활용도", "σmax / σ허용"], ["SF", "안전계수 (하중 배율)"], ["슬링각", "와이어와 수평면이 이루는 각"],
             ["수직반력", "장력 × sin(슬링각)"], ["COG", "무게중심"]], widths=[3, 9])
    chap("C", "부록 C. 해석 로그 요약")
    if f.get("available"):
        s.kv([("FATAL", str(f.get("fatalCount", 0))), ("ERROR", str(f.get("errorCount", 0)))])
        for smp in (f.get("fatalSamples") or []) + (f.get("errorSamples") or []):
            s.note(smp)
    else:
        none_note("F06")
    for w in d.warnings:
        s.note("※ " + w)

    if pages is not None:
        s.toc_fill([(text, pages.get(text, 0), lvl) for text, lvl in entries])
    total = s.finish()
    return sec_pages, total


def _data_sheets(wb: Workbook, d: ReportData):
    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(name=FONT, bold=True)
            cell.fill = PatternFill("solid", fgColor=LIGHT)
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 14

    sheet("Members", ["elementId", "pid", "section", "maxStressMPa", "utilization", "exceeds"],
          ((x.element_id, x.pid, x.section_label, x.stress_mpa, x.utilization, x.exceeds) for x in d.results.all_members))
    sheet("Displacements", ["nodeId", "t1", "t2", "t3", "magnitude"],
          ((x.node_id, x.t1, x.t2, x.t3, x.magnitude) for x in d.results.all_displacements))
    sheet("Wires", ["wireElementId", "groupId", "lugNodeId", "axialForceN", "tensionTon", "angleDeg", "verticalTon",
                    "needsJig", "isCompression", "isSlack"],
          ((x.wire_element_id, x.group_id, x.lug_node_id, x.axial_force_n, x.tension_ton, x.angle_deg, x.vertical_ton,
            x.needs_jig, x.is_compression, x.is_slack) for x in d.results.wires))
