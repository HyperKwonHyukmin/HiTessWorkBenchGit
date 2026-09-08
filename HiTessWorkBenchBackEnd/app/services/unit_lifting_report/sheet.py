"""보고서 레이아웃 프리미티브 — 사내 서식처럼 '고정 페이지 틀' 안에 내용을 채운다.

원본 xlsx 서식은 한 페이지가 (머리글 블록 + 본문 + 바닥글) 로 이뤄진 테두리 상자였고,
그 블록이 페이지마다 반복됐다. 여기서도 같은 구조를 만든다:

    ┌───────────────────────────────────────────────┐  ← 페이지 외곽선(medium)
    │ [로고]        제목(굵게)            부서/문서번호 │  머리글 3행
    │ HULL NO. xxxx │ UNIT NO. xxxxx │ 권상 방식      │  머리글 1행
    ├───────────────────────────────────────────────┤
    │ 본문 50행 (문단·표·그림)                        │
    ├───────────────────────────────────────────────┤
    │ Hi-TESS WorkBench 자동 생성        Page n / N   │  바닥글
    └───────────────────────────────────────────────┘

핵심 규약 — **모든 행 높이가 ROW_PT 로 같다.** 그래야 '행 수 = 세로 공간' 이 되어
페이지 경계를 정확히 계산하고 테두리를 그릴 수 있다. 여러 줄 텍스트는 행을 세로 병합해 쓴다.
본문이 넘치면 `_ensure()` 가 페이지를 닫고 새 페이지를 연다(표는 머리행을 다시 그린다).
"""
from __future__ import annotations

import io
import unicodedata
from typing import Iterable, Optional, Sequence

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

COLS = 12
COL_WIDTH = 7.2
ROW_PT = 13.5          # = 18px 정확 (14pt 는 18.67px → 19px 로 반올림돼 프레임 높이가 틀어진다)
PX_PER_ROW = ROW_PT * 96 / 72      # 행 하나의 픽셀 높이(그림 배치용)

HEADER_ROWS = 5                     # 로고·제목 3행 + HULL/UNIT 1행 + 여백 1행
BODY_ROWS = 50
FOOTER_ROWS = 2
PAGE_ROWS = HEADER_ROWS + BODY_ROWS + FOOTER_ROWS   # 57행 × 13.5pt = 769.5pt < A4 인쇄 높이(781pt)

FONT = "맑은 고딕"
NAVY, GRAY, LIGHT, LINE = "002554", "5B6975", "F3F5F7", "AEB8C2"
STATUS_FILL = {"ok": "E3F6EA", "pass": "E3F6EA", "warn": "FFF3D6", "ng": "FFE0E3", "fail": "FFE0E3",
               "skip": "EEF1F4", "판정 불가": "EEF1F4", "필요": "FFF3D6", "불요": "E3F6EA"}

_THIN = Side(style="thin", color=LINE)
_MEDIUM = Side(style="medium", color=NAVY)
CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def text_width(text) -> int:
    """한글·전각을 2, 나머지를 1 로 센 표시 폭."""
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in str(text))


class ReportSheet:
    def __init__(self, ws, doc_title: str = "", hull: str = "", unit: str = "",
                 method: str = "", department: str = "", drawing_no: str = "",
                 logo_png: Optional[bytes] = None, total_pages: Optional[int] = None):
        self.ws = ws
        ws.title = "Report"
        self.doc_title = doc_title
        self.hull, self.unit, self.method = hull, unit, method
        self.department, self.drawing_no = department, drawing_no
        self.logo = logo_png
        self.total_pages = total_pages

        self.page = 1
        self.row = 1                 # 다음 페이지 프레임이 시작할 행
        self.body_row = 1
        self.body_end = 0
        self.page_start = 1
        self.chapter_pages: dict[str, int] = {}
        self._fig_no = 0
        self._tbl_no = 0
        self._toc_rows: list[int] = []
        self._page_open = False

        for c in range(1, COLS + 1):
            ws.column_dimensions[get_column_letter(c)].width = COL_WIDTH
        ws.sheet_view.showGridLines = False
        ps = ws.page_setup
        ps.orientation = "portrait"
        ps.paperSize = ws.PAPERSIZE_A4
        # ⚠ '페이지에 맞춤(fitToPage)' 을 켜면 Excel 이 **수동 페이지 나누기를 무시**해서
        # 프레임 두 개가 한 장에 겹쳐 인쇄된다. 배율을 100% 로 고정해 프레임 1개 = 1페이지를 보장한다.
        # (12열 × 7.2 ≈ 6.9in < 인쇄폭 7.47in, 57행 × 13.5pt ≈ 10.7in < 인쇄높이 10.85in 이라 잘리지 않는다.)
        ps.scale = 100
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_margins.left = ws.page_margins.right = 0.4
        # 프레임 769.5pt < 인쇄 높이 781pt (여유 11.5pt). 여유가 프레임보다 크면 페이지가 쪼개지고,
        # 0 이면 반올림 한 번에 쪼개진다. 이 여유에 다음 페이지 내용이 비치지 않도록 로고는 첫 행이 아닌
        # 둘째 행에 앵커한다(_open_page 참조).
        ws.page_margins.top = ws.page_margins.bottom = 0.42
        ws.page_margins.header = ws.page_margins.footer = 0.2

    # ── 셀/테두리 저수준 ────────────────────────────────────────────────────
    def _cell(self, row, c0, c1=None, value=None, size=9.5, bold=False, color="18242F",
              align="left", valign="center", wrap=True, fill=None, rows=1, border=False, indent=0):
        c1 = c1 or c0
        r1 = row + rows - 1
        if c1 > c0 or r1 > row:
            self.ws.merge_cells(start_row=row, start_column=c0, end_row=r1, end_column=c1)
        cell = self.ws.cell(row=row, column=c0)
        cell.value = value
        cell.font = Font(name=FONT, size=size, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap, indent=indent)
        if fill:
            for r in range(row, r1 + 1):
                for c in range(c0, c1 + 1):
                    self.ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=fill)
        if border:
            for r in range(row, r1 + 1):
                for c in range(c0, c1 + 1):
                    self.ws.cell(row=r, column=c).border = CELL_BORDER
        return cell

    def _edge(self, row, col, top=None, bottom=None, left=None, right=None):
        """기존 테두리를 유지하며 지정한 변만 교체한다."""
        b = self.ws.cell(row=row, column=col).border
        self.ws.cell(row=row, column=col).border = Border(
            top=top or b.top, bottom=bottom or b.bottom, left=left or b.left, right=right or b.right)

    def _outline(self, r0, r1, side=_MEDIUM, c0=1, c1=COLS):
        for c in range(c0, c1 + 1):
            self._edge(r0, c, top=side)
            self._edge(r1, c, bottom=side)
        for r in range(r0, r1 + 1):
            self._edge(r, c0, left=side)
            self._edge(r, c1, right=side)

    def _hline(self, row, side=_MEDIUM, c0=1, c1=COLS):
        for c in range(c0, c1 + 1):
            self._edge(row, c, bottom=side)

    # ── 페이지 프레임 ───────────────────────────────────────────────────────
    def _open_page(self, cover: bool = False):
        r = self.row
        self.page_start = r
        if r > 1:
            self.ws.row_breaks.append(Break(id=r - 1))
        for i in range(PAGE_ROWS):
            self.ws.row_dimensions[r + i].height = ROW_PT

        if cover:
            # 표지: 머리글 자리를 로고와 문서 종류로만 채운다(큰 제목은 본문에서).
            if self.logo:
                img = XLImage(io.BytesIO(self.logo))
                img.width, img.height = 218, 36
                img.anchor = f"A{r + 1}"
                self.ws.add_image(img)
            self._cell(r, 1, 4, None, rows=3)
            self._cell(r, 5, COLS, "STRUCTURAL REVIEW REPORT", size=9, bold=True, color=GRAY,
                       align="right", rows=3)
            self._cell(r + 3, 1, COLS, None)
        else:
            if self.logo:
                img = XLImage(io.BytesIO(self.logo))
                img.width, img.height = 145, 24
                img.anchor = f"A{r + 1}"
                self.ws.add_image(img)
            self._cell(r, 1, 3, None, rows=3)
            self._cell(r, 4, 9, self.doc_title, size=12.5, bold=True, color=NAVY, align="center", rows=3)
            right = self.department or (f"DWG. {self.drawing_no}" if self.drawing_no and self.drawing_no != "-" else "")
            self._cell(r, 10, COLS, right, size=8, color=GRAY, align="right", wrap=False, rows=3)
            # HULL / UNIT / 권상 방식 — 원본 서식의 O4·Y4·AI4 자리
            self._cell(r + 3, 1, 4, f"HULL NO. {self.hull}", size=9, bold=True, align="center",
                       fill=LIGHT, border=True)
            self._cell(r + 3, 5, 8, f"UNIT NO. {self.unit}", size=9, bold=True, align="center",
                       fill=LIGHT, border=True)
            self._cell(r + 3, 9, COLS, self.method, size=9, bold=True, align="center",
                       fill=LIGHT, border=True)
        self._hline(r + 3)
        self.body_row = r + HEADER_ROWS
        self.body_end = r + HEADER_ROWS + BODY_ROWS - 1
        self._page_open = True

    def _close_page(self):
        if not self._page_open:
            return
        r = self.page_start
        foot = r + HEADER_ROWS + BODY_ROWS          # 바닥글 첫 행
        self._hline(foot - 1)
        total = f" / {self.total_pages}" if self.total_pages else ""
        self._cell(foot + 1, 1, 8, "본 보고서는 Hi-TESS WorkBench 가 해석 결과로부터 자동 생성했습니다.",
                   size=8, color=GRAY)
        self._cell(foot + 1, 9, COLS, f"Page {self.page}{total}", size=8, color=GRAY, align="right")
        self._outline(r, r + PAGE_ROWS - 1)
        self.row = r + PAGE_ROWS
        self.page += 1
        self._page_open = False

    def _ensure(self, rows: int, cover: bool = False):
        """본문에 rows 행이 들어갈 자리를 확보한다. 모자라면 페이지를 넘긴다."""
        if not self._page_open:
            self._open_page(cover=cover)
        if self.body_row + rows - 1 > self.body_end:
            self._close_page()
            self._open_page()
            return True
        return False

    def _take(self, rows: int) -> int:
        row = self.body_row
        self.body_row += rows
        return row

    def new_page(self):
        self._close_page()
        self._open_page()

    def finish(self) -> int:
        """마지막 페이지를 닫고 인쇄 영역을 확정한다. 총 페이지 수를 돌려준다."""
        self._close_page()
        self.ws.print_area = f"A1:{get_column_letter(COLS)}{self.row - 1}"
        return self.page - 1

    # ── 행 수 계산 ──────────────────────────────────────────────────────────
    @staticmethod
    def _rows_for(text, cols: int = COLS, size: float = 9.5) -> int:
        cap = max(10, int(cols * COL_WIDTH * (11.0 / size)))
        n = 0
        for part in str(text).split("\n"):
            n += max(1, -(-text_width(part) // cap))
        return n

    # ── 공개 프리미티브 ─────────────────────────────────────────────────────
    def cover(self, title: str, facts: Sequence[tuple[str, str]], verdicts: Sequence[tuple[str, str, str]]):
        self._ensure(1, cover=True)
        r = self._take(4)
        self._cell(r, 1, COLS, title, size=22, bold=True, color=NAVY, align="center", valign="center", rows=4)
        r = self._take(2)
        self._cell(r, 1, COLS, "Structural Review Report for Module Unit Lifting",
                   size=10.5, color=GRAY, align="center", rows=2)
        self._take(3)
        for k, v in facts:
            r = self._take(2)
            self._cell(r, 3, 5, k, size=10, bold=True, color=GRAY, align="right", rows=2)
            self._cell(r, 6, 10, v, size=10.5, rows=2)
        self._take(3)
        r = self._take(1)
        self._cell(r, 1, COLS, "핵심 판정", size=10, bold=True, color=GRAY, align="center")
        span = COLS // max(1, len(verdicts))
        r = self._take(1)
        c = 1
        for label, _value, _status in verdicts:
            self._cell(r, c, c + span - 1, label, size=9, color=GRAY, align="center", fill=LIGHT, border=True)
            c += span
        r = self._take(3)
        c = 1
        for _label, value, status in verdicts:
            self._cell(r, c, c + span - 1, value, size=18, bold=True, align="center", rows=3,
                       fill=STATUS_FILL.get(str(status).lower(), LIGHT), border=True)
            c += span
        self.new_page()

    def toc_placeholder(self, n_entries: int):
        if not self._page_open:
            self._open_page()
        elif self.body_row != self.page_start + HEADER_ROWS:
            self.new_page()
        r = self._take(2)
        self._cell(r, 1, COLS, "목차 (Contents)", size=14, bold=True, color=NAVY, rows=2)
        self._toc_rows = []
        for _ in range(n_entries):
            self._ensure(1)
            self._toc_rows.append(self._take(1))
        self.new_page()

    def toc_fill(self, entries: Sequence[tuple]):
        """entries: (text, page) 또는 (text, page, level). level 0 = 장(굵게), 1 = 절(들여쓰기)."""
        for row, entry in zip(self._toc_rows, entries):
            text, page = entry[0], entry[1]
            level = entry[2] if len(entry) > 2 else 0
            self._cell(row, 1, 10, text, size=9.5, bold=(level == 0), indent=level * 2)
            self._cell(row, 11, COLS, page, size=9.5, align="right")

    def chapter(self, key: str, text: str) -> int:
        # 장은 항상 새 페이지에서 시작한다. 단 이미 빈 페이지의 맨 위면 그대로 쓴다.
        if not self._page_open:
            self._open_page()
        elif self.body_row != self.page_start + HEADER_ROWS:
            self.new_page()
        self.chapter_pages[key] = self.page
        r = self._take(2)
        self._cell(r, 1, COLS, text, size=15, bold=True, color=NAVY, rows=2)
        self._hline(r + 1, side=Side(style="medium", color=NAVY))
        self._take(1)
        return self.page

    def section(self, text: str) -> int:
        self._ensure(3)
        self._take(1)
        r = self._take(2)
        self._cell(r, 1, COLS, text, size=11.5, bold=True, color=NAVY, rows=2)
        return self.page

    def subsection(self, text: str):
        self._ensure(2)
        r = self._take(2)
        self._cell(r, 1, COLS, text, size=10, bold=True, rows=2)

    def para(self, text: str, size: float = 9.5, color: str = "18242F"):
        n = self._rows_for(text, COLS, size)
        self._ensure(n)
        r = self._take(n)
        self._cell(r, 1, COLS, text, size=size, color=color, valign="top", rows=n)

    def bullet(self, text: str):
        n = self._rows_for(text, COLS - 1, 9.5)
        self._ensure(n)
        r = self._take(n)
        self._cell(r, 1, 1, "•", align="center", valign="top", rows=n)
        self._cell(r, 2, COLS, text, valign="top", rows=n)

    def note(self, text: str):
        self.para(text, size=8.5, color=GRAY)

    def kv(self, rows: Sequence[tuple[str, str]], key_cols: int = 4):
        for k, v in rows:
            n = max(1, self._rows_for(v, COLS - key_cols, 9.5))
            self._ensure(n)
            r = self._take(n)
            self._cell(r, 1, key_cols, k, size=9, bold=True, fill=LIGHT, border=True, rows=n)
            self._cell(r, key_cols + 1, COLS, v, size=9.5, border=True, valign="center", rows=n)
        self._ensure(1)
        self._take(1)

    def _table_header(self, headers, widths):
        self._ensure(1)
        r = self._take(1)
        c = 1
        for h, w in zip(headers, widths):
            self._cell(r, c, c + w - 1, h, size=9, bold=True, align="center", fill=LIGHT, border=True, wrap=False)
            c += w

    def table(self, headers: Sequence[str], rows: Iterable[Sequence], widths: Sequence[int],
              align: Optional[Sequence[str]] = None, status_col: Optional[int] = None,
              number_format: str = "General", title: Optional[str] = None) -> int:
        """widths 합계는 COLS 이하. 페이지가 넘어가면 머리행을 다시 그린다."""
        assert sum(widths) <= COLS, f"표 폭 초과: {widths}"
        if title:
            self._tbl_no += 1
            self._ensure(4)                       # 표 제목만 페이지 끝에 남지 않게
            r = self._take(1)
            self._cell(r, 1, COLS, f"표 {self._tbl_no}. {title}", size=9, bold=True, color=GRAY)
        self._table_header(headers, widths)
        n = 0
        for data in rows:
            if self._ensure(1):
                self._table_header(headers, widths)   # 새 페이지에 머리행 반복
            r = self._take(1)
            c = 1
            for i, (v, w) in enumerate(zip(data, widths)):
                a = (align[i] if align else ("right" if isinstance(v, (int, float)) and not isinstance(v, bool) else "left"))
                fill = STATUS_FILL.get(str(v).lower()) if (status_col is not None and i == status_col) else None
                cell = self._cell(r, c, c + w - 1, v, size=9, align=a, wrap=False, fill=fill, border=True)
                if isinstance(v, float):
                    cell.number_format = number_format
                c += w
            n += 1
        if n == 0:
            self._ensure(1)
            r = self._take(1)
            self._cell(r, 1, COLS, "(해당 없음)", size=9, color=GRAY, align="center", border=True)
        self._ensure(1)
        self._take(1)
        return self._tbl_no

    # 그림 기본 크기(px). 12열 × 55px = 660px 이 프레임 폭이므로 그 안에 들어와야 한다
    # (넘기면 컬러바·축라벨이 테두리를 뚫는다).
    FIG_W, FIG_H = 640, 410
    FIG_MIN_ROWS = 20                   # 이보다 좁은 잔여 공간이면 줄이지 말고 다음 페이지로

    def figure(self, png: bytes, caption: str, width_px: int = 0, height_px: int = 0) -> int:
        """그림을 넣는다. 남은 공간이 모자라면 비율을 유지한 채 줄여 그 페이지에 담고,
        그마저 좁으면 다음 페이지로 넘긴다(그림 하나가 페이지를 통째로 비우지 않게)."""
        self._fig_no += 1
        width_px = width_px or self.FIG_W
        height_px = height_px or self.FIG_H
        if not self._page_open:
            self._open_page()
        want = -(-height_px // int(PX_PER_ROW))
        free = self.body_end - self.body_row + 1 - 2          # 캡션 + 여백 1행씩
        if free < min(want, self.FIG_MIN_ROWS):
            self.new_page()
            free = BODY_ROWS - 2
        rows = min(want, free, BODY_ROWS - 2)
        if rows < want:                                        # 비율 유지 축소
            scale = (rows * PX_PER_ROW) / height_px
            width_px, height_px = int(width_px * scale), int(rows * PX_PER_ROW)
        r = self._take(rows)
        img = XLImage(io.BytesIO(png))
        img.width, img.height = width_px, height_px
        img.anchor = f"A{r}"
        self.ws.add_image(img)
        cr = self._take(1)
        self._cell(cr, 1, COLS, f"그림 {self._fig_no}. {caption}", size=8.5, color=GRAY, align="center")
        self._take(1)
        return self._fig_no

    def placeholder(self, caption: str, reason: str):
        """그림 렌더 실패 자리표시."""
        self._fig_no += 1
        self._ensure(5)
        r = self._take(3)
        self._cell(r, 1, COLS, f"[그림 {self._fig_no} 렌더 실패] {reason}", size=9, color="B00020",
                   align="center", fill=LIGHT, border=True, rows=3)
        cr = self._take(1)
        self._cell(cr, 1, COLS, f"그림 {self._fig_no}. {caption}", size=8.5, color=GRAY, align="center")
        self._take(1)
