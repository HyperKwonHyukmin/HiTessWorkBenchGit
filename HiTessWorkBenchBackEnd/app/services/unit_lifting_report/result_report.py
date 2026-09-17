"""사내 표준 서식(ver2609) 기반 '결과 레포트' 생성.

`ReportFormat/ver02/260907_Unit_Lifting_Report_ver2609.xlsx` 를 그대로 서식으로 쓰고
해석 결과와 3D 캡처만 채운다. 서식의 셀 배치·수식·인쇄 설정은 건드리지 않는다.

  페이지 1 (행 1–54)   : 제목 · 주의사항(5줄) · 요약표 · 권상 위치 3D(평면/등각)
  페이지 2 (행 55–108) : 1. 구조 해석 결과 — 변위·응력 3D + Hook/Trolley 반력표
  페이지 3 (행 109–162): 2. 서포트 배치 — 가서포트가 없으면 통째로 삭제한다(사용자 결정)

서식 원본은 사내 DRM 대상(.xlsx)이라 `.bin` 사본을 배포본으로 둔다. [[reference_drm_xlsx_template_pattern]]
"""
from __future__ import annotations

import copy
import io
import math
import os
import re
from datetime import datetime
from typing import Any, Optional

from openpyxl import load_workbook

from .collector import ReportOptions, collect, format_contact
from . import figures3d

TEMPLATE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "templates"))
TEMPLATE_BIN = os.path.join(TEMPLATE_DIR, "unit_lifting_result_template.bin")
TEMPLATE_XLSX = os.path.join(TEMPLATE_DIR, "unit_lifting_result_template.xlsx")

# 서식의 캡처 자리(이미지 3~8) 순서. 앞 3개(0~2)는 HD 로고라 보존한다.
IMAGE_SLOTS = ("modelTop", "modelHome", "displacement", "stress", "supportTop", "supportHome")
LOGO_COUNT = 3
SUPPORT_PAGE_FIRST_ROW = 109
SUPPORT_PAGE_LAST_ROW = 162
SUPPORT_PAGE_BREAK_ID = 108
TON_FORCE_N = 9800.0

# Hook/Trolley 표: 그룹 블록 시작 행과 블록당 wire 행 수
HOOK_BLOCK_ROWS = (92, 96, 100)
WIRES_PER_BLOCK = 4

# 주의 사항 — 서식의 F7~F11(F~BA 병합) 5줄. 1~3행은 사내 서식이 갖고 있는 문구라 건드리지 않고,
# 4행(고정)·5행(사용자 입력)만 코드가 채운다. 문구를 고칠 일이 생기면 여기만 고치면 된다.
# ⚠ 서식(.bin)은 이 5줄 자리를 만들려고 2026-09-17 에 1페이지를 재배치했다 — 주의 사항이 2줄
#   늘어난 만큼 요약표·각주가 2행 내려갔고(11~14행 → 13~16행), 3D 그림 두 장은 각주 아래와
#   바닥글 위의 빈 줄을 하나씩 내주고 같은 크기(17행)로 아래에 붙었다. 페이지 행 수(1~54)와
#   페이지 나누기는 그대로다. 사내 서식이 새로 배포되면 이 재배치를 다시 해야 한다.
NOTICE_ADDED_ROW = 10
NOTICE_EXTRA_ROW = 11
NOTICE_ADDED_LINE = "- 200mm(현장 기준 확인) 현장 권상 및 안전성 확인 후 최종 체결 위치 선정"


def generate_result_report(result_info: dict[str, Any], options: dict[str, Any] | None = None, *,
                           generated_by: str = "", generator: dict[str, Any] | None = None,
                           ) -> tuple[str, bytes, list[str], dict[str, Any]]:
    """``(파일명, xlsx 바이트, 경고, 요약)``. 디스크에 쓰지 않는다."""
    opts = ReportOptions.from_payload(options)
    if not opts.author:
        opts.author = generated_by or "-"
    opts.contact = opts.contact or format_contact(generator, generated_by)   # 바닥글 '문의' 줄
    now = datetime.now()
    data = collect(result_info, opts, generated_at=now.strftime("%Y-%m-%d %H:%M"))
    warnings: list[str] = list(data.warnings)

    figs, fig_errors = figures3d.render_result_figures(data)
    for key, why in fig_errors.items():
        warnings.append(f"3D 캡처 '{key}' 렌더 실패: {why}")

    wb = _load_template()
    ws = wb["Report"]
    # ⚠ 서식 원본에는 print_title_rows='$1:$3' (1~3행을 매 페이지 상단에 반복 인쇄) 가 남아 있다.
    # 이 서식은 페이지마다 자체 머리글(행 57·111)을 이미 갖고 있어서, 2·3페이지 맨 위에
    # 1페이지 머리글이 한 번 더 찍혀 겹친다(원본 서식을 그대로 인쇄해도 재현된다). 내용은
    # 건드리지 않고 이 반복 인쇄 설정만 끈다.
    # openpyxl 의 print_title_rows setter 는 None 을 무시하므로 내부 속성을 직접 비운다.
    ws._print_rows = None
    ws.print_title_rows = None
    has_support = bool(data.edits and data.edits.support_beams)
    _fill_titles(ws, data, has_support)
    _fill_notice(ws, opts)
    _fill_summary(ws, data)
    _fill_hook_table(ws, data, warnings)
    _replace_images(ws, figs, has_support, warnings)
    if not has_support:
        _drop_support_page(ws)
    _fit_form_text(ws)
    _solid_frame(ws)
    _print_setup(ws)

    buf = io.BytesIO()
    wb.save(buf)

    i = data.identity
    name = (f"{_safe(i.hull_no)}-{_safe(i.unit_no)}_{_safe(i.lifting_method)}_"
            f"{_safe(i.title_prefix)}_권상_구조_해석_보고서_{now:%Y%m%d}.xlsx")
    summary = {
        "totalMassTon": data.model.total_mass_ton,
        "yieldStrengthMPa": data.results.yield_mpa,
        "allowableMPa": data.results.allowable_mpa,
        "maxDisplacementMm": data.results.max_displacement_mm,
        "maxStressMPa": data.results.max_stress_mpa,
        "structure": data.verdicts.structure,
        "jigRequired": data.verdicts.jig_required,
        "supportPage": has_support,
        "figureCount": len(figs),
        "warningCount": len(warnings),
    }
    return name, buf.getvalue(), warnings, summary


# ── 서식 로드 ────────────────────────────────────────────────────────────────
def _load_template():
    """DRM 상태와 무관하게 열리는 서식을 고른다(.bin 우선, 없으면 .xlsx)."""
    tried: list[str] = []
    for path in (TEMPLATE_BIN, TEMPLATE_XLSX):
        if not os.path.isfile(path):
            continue
        with open(path, "rb") as fh:
            raw = fh.read()
        if raw[:2] != b"PK":
            tried.append(f"{os.path.basename(path)}(DRM 암호화 상태)")
            continue
        return load_workbook(io.BytesIO(raw), keep_links=False)
    detail = f" — 시도: {', '.join(tried)}" if tried else ""
    raise FileNotFoundError(f"결과 레포트 서식을 열 수 없습니다: {TEMPLATE_BIN}{detail}")


# ── 셀 채우기 ────────────────────────────────────────────────────────────────
def _fill_titles(ws, d, has_support: bool):
    """페이지마다 반복되는 제목·호선·유닛·권상방식."""
    title = f"{d.identity.title_prefix} 권상 구조 해석 보고서"
    rows = [(3, 4), (57, 58)] + ([(111, 112)] if has_support else [])
    for title_row, info_row in rows:
        ws[f"O{title_row}"] = title
        ws[f"O{info_row}"] = f"HULL NO. {d.identity.hull_no}"
        ws[f"Y{info_row}"] = f"UNIT NO. {d.identity.unit_no}"
        ws[f"AI{info_row}"] = d.identity.lifting_method
    # 문의처는 보고서를 만든 WorkBench 사용자(이름/직급/부서) — 표지의 '작성자' 입력과 별개다
    footer = ("본 보고서는 Hi-TESS WorkBench를 통해 자동 생성되었습니다.\n"
              f"문의 | {d.identity.contact}   생성일 {d.identity.generated_at}")
    for row in ([52, 106] + ([160] if has_support else [])):
        ws[f"F{row}"] = footer


def _fill_notice(ws, opts):
    """주의 사항 4행(고정 문구)·5행(사용자가 모달에 적은 한 줄)을 채운다.

    5행이 비면 그 줄을 **감추고** 박스 아랫변을 4행으로 올린다(사용자 결정 2026-09-17) —
    빈 줄이 남으면 서식에 빈 칸이 한 줄 뚫린 것처럼 보인다.
    """
    ws[f"F{NOTICE_ADDED_ROW}"] = NOTICE_ADDED_LINE
    extra = (opts.extra_notice or "").strip()
    if extra:
        # 다른 줄과 같은 글머리표를 붙인다. `_fit_form_text` 도 '-' 로 시작하는 F열 병합 칸만
        # '셀에 맞춤' 대상으로 잡으므로, 이 접두사가 있어야 긴 문구가 잘리지 않고 줄어든다.
        ws[f"F{NOTICE_EXTRA_ROW}"] = extra if extra.startswith("-") else f"- {extra}"
        return
    ws.row_dimensions[NOTICE_EXTRA_ROW].hidden = True
    ws.row_dimensions[NOTICE_EXTRA_ROW].height = 0
    _move_bottom_border(ws, NOTICE_EXTRA_ROW, NOTICE_ADDED_ROW)


def _move_bottom_border(ws, src_row: int, dst_row: int):
    """`src_row` 의 아래 테두리를 `dst_row` 로 옮긴다(감춘 행의 테두리는 인쇄되지 않는다)."""
    from openpyxl.styles import Border

    for col in range(1, ws.max_column + 1):
        bottom = ws.cell(src_row, col).border.bottom
        cell = ws.cell(dst_row, col)
        b = cell.border
        cell.border = Border(left=b.left, right=b.right, top=b.top, bottom=copy.copy(bottom),
                             diagonal=b.diagonal, diagonalUp=b.diagonalUp, diagonalDown=b.diagonalDown)


def _fill_summary(ws, d):
    """요약표 — 허용응력(V15)·검토결과(AT15)는 서식의 수식을 그대로 둔다."""
    r = d.results
    ws["F15"] = round(d.model.total_mass_ton, 3)
    ws["N15"] = round(r.yield_mpa, 1)
    ws["AD15"] = round(r.max_displacement_mm, 2)
    ws["AL15"] = round(r.max_stress_mpa, 2)
    ws["N74"] = f"[ 최대 변형 결과 : {r.max_displacement_mm:.4g} mm ]"
    ws["N88"] = f"[ 최대 응력 결과 : {r.max_stress_mpa:.4g} MPa ]"


def _fill_hook_table(ws, d, warnings: list[str]):
    """Hook/Trolley 별 반력·장력. 서식은 3그룹 × 4 wire 까지만 표시한다."""
    ws["AD92"] = d.options.jig_limit_ton                  # 안전 하중 기준(기본 6.2 ton)
    for start in HOOK_BLOCK_ROWS:                          # 서식 초기값 비우기
        ws[f"F{start}"] = None
        ws[f"R{start}"] = None
    for row in range(HOOK_BLOCK_ROWS[0], HOOK_BLOCK_ROWS[-1] + WIRES_PER_BLOCK):
        ws[f"X{row}"] = None

    by_group: dict[int, list] = {}
    for w in d.results.wires:
        by_group.setdefault(w.group_id, []).append(w)
    group_ids = sorted(by_group)
    if len(group_ids) > len(HOOK_BLOCK_ROWS):
        warnings.append(f"서식은 Hook/Trolley {len(HOOK_BLOCK_ROWS)}개까지 표시합니다. "
                        f"{len(group_ids)}개 중 앞의 {len(HOOK_BLOCK_ROWS)}개만 기록했습니다.")

    for block, start in enumerate(HOOK_BLOCK_ROWS):
        for offset in range(WIRES_PER_BLOCK):
            ws[f"L{start + offset}"] = None
        if block >= len(group_ids):
            ws[f"F{start}"] = "-"
            for offset in range(WIRES_PER_BLOCK):
                ws[f"AJ{start + offset}"] = None
            continue
        gid = group_ids[block]
        wires = sorted(by_group[gid], key=lambda w: (w.lug_node_id, w.wire_element_id))
        # 그룹·러그를 캡처 그림의 라벨(G1·N34)과 같은 표기로 적어 어느 와이어인지 그림에서 찾을 수 있게 한다
        ws[f"F{start}"] = f"G{gid}"
        if len(wires) > WIRES_PER_BLOCK:
            warnings.append(f"그룹 {gid}의 wire {len(wires)}개 중 {WIRES_PER_BLOCK}개만 서식에 기록했습니다.")
        for offset in range(len(wires), WIRES_PER_BLOCK):
            ws[f"AJ{start + offset}"] = None          # 없는 wire 행에 '지그 불요' 가 뜨지 않게
        reaction = 0.0
        has_reaction = False
        for offset, w in enumerate(wires[:WIRES_PER_BLOCK]):
            ws[f"L{start + offset}"] = f"G{gid}-N{w.lug_node_id}"
            ws[f"X{start + offset}"] = round(w.tension_ton, 3)
            if w.vertical_ton is not None:
                reaction += w.vertical_ton
                has_reaction = True
            else:
                warnings.append(f"그룹 {gid}, Lug {w.lug_node_id}의 슬링각이 없어 반력을 기록하지 못했습니다.")
        if has_reaction:
            ws[f"R{start}"] = round(reaction, 3)


# ── 이미지 ───────────────────────────────────────────────────────────────────
# 서식의 자리표시 그림은 TwoCellAnchor(시작 셀→끝 셀) 라 Excel 이 그림을 **셀 박스에 맞춰 늘린다.**
# 그런데 이 서식은 1.71자짜리 아주 좁은 열 56개로 짜여 있어, Excel 인쇄 엔진이 열 너비를
# 화면(정수 px)과 다르게(소수 px) 계산하면서 셀 박스가 인쇄에서 ~8% 더 넓어진다. 그 결과 PDF 의
# 3D 그림이 옆으로 11% 늘어나 보였다(실측: 원본 1.731 → 표시 1.932). 그림을 **절대 크기**
# (OneCellAnchor + ext) 로 붙이면 화면·인쇄 어디서도 늘어나지 않는다. 크기는 셀 박스 안에
# 비율을 지켜 담고, 남는 폭은 양쪽으로 나눠 가운데에 둔다.
#
# 셀 박스 px 는 Excel 규칙으로 계산한다 — 열: int((256·w + int(128/MDW))/256 · MDW), MDW 는
# 통합 문서 기본 글꼴(Calibri 11)의 숫자 폭 7px(96dpi). 행: pt × 96/72 반올림.
_MDW_PX = 7            # Calibri 11 @96dpi — 서식의 Normal 스타일 글꼴(실측 col 1.71자 = 12px = 9.0pt)
_DEFAULT_COL_WIDTH = 8.43
_EMU_PER_PX = 9525


def _col_width_chars(ws, col_idx0: int) -> float:
    """0-based 열 번호의 너비(문자 단위). openpyxl 은 열 범위(min~max)로 저장하므로 직접 찾는다.
    ⚠ ws.column_dimensions[letter] 로 읽으면 없는 열에 기본 13 짜리 항목이 새로 생겨 틀린다."""
    idx = col_idx0 + 1
    for dim in ws.column_dimensions.values():
        lo, hi = dim.min or 0, dim.max or 0
        if lo <= idx <= hi and dim.width:
            return float(dim.width)
    return float(ws.sheet_format.defaultColWidth or _DEFAULT_COL_WIDTH)


def _col_px(ws, col_idx0: int) -> int:
    w = _col_width_chars(ws, col_idx0)
    return int((256 * w + int(128 / _MDW_PX)) / 256 * _MDW_PX)


def _row_px(ws, row_idx0: int) -> int:
    pt = ws.row_dimensions[row_idx0 + 1].height
    if pt is None:
        pt = ws.sheet_format.defaultRowHeight or 15.0
    return int(round(float(pt) * 96 / 72))


def _cell_box_px(ws, anchor) -> tuple[int, int]:
    """TwoCellAnchor 가 덮는 셀 박스 크기(px)."""
    f, t = anchor._from, anchor.to
    w = sum(_col_px(ws, c) for c in range(f.col, t.col)) + (t.colOff - f.colOff) // _EMU_PER_PX
    h = sum(_row_px(ws, r) for r in range(f.row, t.row)) + (t.rowOff - f.rowOff) // _EMU_PER_PX
    return max(1, int(w)), max(1, int(h))


def _absolute_anchor(ws, placeholder_anchor, img_w: int, img_h: int):
    """자리표시 앵커 → 같은 자리에 그림을 절대 크기로 놓는 OneCellAnchor(가운데 정렬).
    이미 OneCellAnchor 면 크기만 유지한 채 그대로 쓴다."""
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D

    if not isinstance(placeholder_anchor, TwoCellAnchor):
        return copy.copy(placeholder_anchor), (img_w, img_h)
    box_w, box_h = _cell_box_px(ws, placeholder_anchor)
    w, h = _contain(img_w, img_h, box_w, box_h)
    f = placeholder_anchor._from
    marker = AnchorMarker(col=f.col, colOff=f.colOff + ((box_w - w) // 2) * _EMU_PER_PX,
                          row=f.row, rowOff=f.rowOff + ((box_h - h) // 2) * _EMU_PER_PX)
    return OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(w * _EMU_PER_PX, h * _EMU_PER_PX)), (w, h)


def _replace_images(ws, figs: dict[str, bytes], has_support: bool, warnings: list[str]):
    """서식의 캡처 자리(이미지 3~8)를 새 3D 렌더로 갈아끼운다. 앞 3개 로고는 보존."""
    from openpyxl.drawing.image import Image as XLImage

    existing = list(ws._images)
    if len(existing) < LOGO_COUNT + len(IMAGE_SLOTS):
        raise ValueError(f"서식의 이미지 자리 수가 예상과 다릅니다 "
                         f"(expected >= {LOGO_COUNT + len(IMAGE_SLOTS)}, got {len(existing)}).")
    logos = existing[:LOGO_COUNT]
    if not has_support:
        logos = [im for im in logos if im.anchor._from.row < SUPPORT_PAGE_FIRST_ROW - 1]
    # 1페이지 로고도 TwoCellAnchor 라 같이 늘어난다(실측 1.10배) — 같은 방법으로 고정한다.
    for logo in logos:
        logo.anchor, (logo.width, logo.height) = _absolute_anchor(ws, logo.anchor, logo.width, logo.height)
    placeholders = existing[LOGO_COUNT:LOGO_COUNT + len(IMAGE_SLOTS)]
    ws._images = logos

    slots = IMAGE_SLOTS if has_support else IMAGE_SLOTS[:4]
    for key, placeholder in zip(slots, placeholders):
        png = figs.get(key)
        if png is None:
            warnings.append(f"{key} 캡처가 없어 해당 자리를 비웠습니다.")
            continue
        image = XLImage(io.BytesIO(png))
        image.anchor, (image.width, image.height) = _absolute_anchor(
            ws, placeholder.anchor, image.width, image.height)
        ws.add_image(image)


def _contain(w: float, h: float, max_w: float, max_h: float) -> tuple[int, int]:
    """비율을 유지한 채 서식 자리 안에 들어가게 축소한다(왜곡 없음)."""
    if w <= 0 or h <= 0:
        return (int(max_w), int(max_h))
    scale = min(max_w / w, max_h / h)
    return (max(1, int(w * scale)), max(1, int(h * scale)))


# ── 서식 글자 맞춤 · 인쇄 설정 ──────────────────────────────────────────────
# 서식의 제목·주의사항은 셀 폭에 딱 맞춘 12~16pt 고정 글자라, 문구가 조금만 길어도(예: Side Passage
# 제목, 주의사항 셋째 줄) 셀 밖으로 넘쳐 테두리를 덮고 인쇄 영역 끝에서 잘린다. 줄바꿈 없이
# '셀에 맞춤(shrinkToFit)' 을 켜 넘치는 칸만 글자를 줄인다 — 들어가는 칸은 그대로다.
_FORM_TEXT_COLUMNS = ("O", "Y", "AI", "AS")   # 제목 · HULL · UNIT · 권상방식 · 부서


def _fit_form_text(ws):
    from openpyxl.styles import Alignment

    targets = []
    for rng in ws.merged_cells.ranges:
        cell = ws.cell(rng.min_row, rng.min_col)
        v = cell.value
        if not isinstance(v, str) or not v.strip():
            continue
        col = cell.column_letter
        if col in _FORM_TEXT_COLUMNS or (col == "F" and v.lstrip().startswith("-")):
            targets.append(cell)
    for cell in targets:
        a = cell.alignment
        cell.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical, indent=a.indent,
                                   text_rotation=a.text_rotation, wrap_text=False, shrink_to_fit=True)


# 서식의 좌우 페이지 외곽선이 놓인 자리. 왼쪽은 D|E 열경계, 오른쪽은 BB|BC 열경계이고
# 한 경계를 두 칸(앞 칸의 right, 뒤 칸의 left)이 나눠 갖는다. 같은 선이 구간에 따라
# **열 스타일**(D 열 전체의 right)로도, **셀 테두리**(23~43행 등)로도 찍혀 있어 둘 다 봐야 한다.
_FRAME_EDGES = ((4, "right"), (5, "left"), (54, "right"), (55, "left"))


def _solid_frame(ws):
    """페이지 외곽선의 hairline(0.12pt)을 머리글 블록과 같은 thin(0.96pt)으로 올린다.

    hairline 은 Excel 화면에서는 1px 회색 선이지만 인쇄·PDF 에서는 0.12pt 라 뷰어 배율에 따라
    사라졌다 나타나 **선이 토막나 보인다**(사용자 신고 2026-09-17). 서식의 머리글 블록(3~4행)은
    이미 thin 이라, 한 선의 굵기가 페이지 중간에서 8배 바뀌는 것도 같은 인상을 준다.
    굵기만 올리고 색·위치·그 밖의 스타일(글꼴·채움)은 그대로 둔다.
    """
    from openpyxl.styles import Border, Side
    from openpyxl.utils import get_column_letter

    def solidify(holder) -> bool:
        border = holder.border
        side = getattr(border, edge, None)
        if side is None or side.style != "hair":
            return False
        sides = {k: getattr(border, k) for k in ("left", "right", "top", "bottom")}
        sides[edge] = Side(style="thin", color=side.color)
        holder.border = Border(**sides)
        return True

    for col, edge in _FRAME_EDGES:
        dim = ws.column_dimensions.get(get_column_letter(col))   # 열 전체에 걸린 hairline
        if dim is not None:
            solidify(dim)
        for row in range(1, ws.max_row + 1):                      # 셀에 직접 찍힌 hairline
            solidify(ws.cell(row, col))


def _print_setup(ws):
    """페이지를 좌우 가운데에 놓는다. 원본 서식은 왼쪽 0.59in·오른쪽 0.20in 이라 인쇄물이 왼쪽으로
    치우쳐 보였다. 배율(97%)·세로 여백·페이지 나누기는 서식 그대로 둔다 — 100% 면 1페이지가
    759pt 로 인쇄 높이(757pt)를 2pt 넘겨 페이지가 쪼개진다."""
    # 서식은 왼쪽에 빈 여백 열(A~C ≈ 23pt)이, 오른쪽에는 BC 한 열(≈ 9pt)만 있어 인쇄 영역을 그대로
    # 가운데 맞추면 눈에 보이는 내용이 오른쪽으로 ~12pt 치우친다(실측 좌 72 / 우 48). 그만큼 왼쪽
    # 여백을 줄여 보이는 내용 기준으로 좌우 60pt 씩 맞춘다.
    ws.page_margins.left, ws.page_margins.right = 0.22, 0.56
    ws.print_options.horizontalCentered = True


# ── 서포트 페이지 삭제 ────────────────────────────────────────────────────────
def _drop_support_page(ws):
    """가서포트가 없으면 3페이지를 통째로 지운다(사용자 결정)."""
    ws.delete_rows(SUPPORT_PAGE_FIRST_ROW, SUPPORT_PAGE_LAST_ROW - SUPPORT_PAGE_FIRST_ROW + 1)
    # RowBreak.count 는 읽기 전용이라 brk 리스트만 갈아끼운다(count 는 파생 값).
    ws.row_breaks.brk = [b for b in ws.row_breaks.brk if b.id != SUPPORT_PAGE_BREAK_ID]
    area = ws.print_area[0] if isinstance(ws.print_area, list) and ws.print_area else ws.print_area
    if area:
        ws.print_area = re.sub(r"\$\d+$", f"${SUPPORT_PAGE_FIRST_ROW - 1}", area)


def _safe(v) -> str:
    return re.sub(r'[<>:"/\\|?*\s]+', "_", str(v or "unknown")).strip("_.") or "unknown"
