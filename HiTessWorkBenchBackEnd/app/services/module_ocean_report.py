"""Mooring 보고서 형식을 따르는 Module Unit 해상 운송 XLSX 보고서.

그림은 **백엔드가 그린 PNG**(`module_ocean_figures`)를 받는다. 화면 캡처를 올려
받던 시절에는 해석 화면이 떠 있어야만 보고서가 나왔고 이력에서 다시 뽑을 수도
없었다 — Unit 권상 보고서와 같은 방침으로 바꿨다(CLAUDE.md).
"""
from __future__ import annotations

import io
import json
import os
import re
from datetime import date
from typing import Any, Dict, Optional, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

from .module_ocean_results import MOMENT_CHECK_TOL, REACTION_CHECK_TOL

NAVY, PALE, SLATE = "17365D", "D9EAF7", "27364A"
WHITE, LINE, GREEN, RED = "FFFFFF", "AAB7C4", "D9EAD3", "F4CCCC"
PAGE_ROWS, LAST_COL = 55, 55  # A:BC — Mooring 보고서와 같은 미세 격자 문서 구조
DEFAULT_ROW_PX = 20           # Excel 기본 행 높이 15pt = 20px
PARA_COLS = 45                # _paragraph 병합 폭(6~50 열)
PARA_CHARS = int(PARA_COLS * 1.72)   # 열 폭 1.72 → 한 줄에 들어가는 대략 문자 수


def _read_json(path: str | None) -> dict[str, Any]:
    if not path:
        return {}
    with open(path, "r", encoding="utf-8") as fh:
        value = json.load(fh)
    if not isinstance(value, dict):
        raise ValueError(f"결과 JSON 형식이 올바르지 않습니다: {os.path.basename(path)}")
    return value


def _png(blob: bytes) -> bytes:
    """렌더러가 준 바이트가 정말 PNG 인지만 본다(파일이 아니라 메모리에서 온다)."""
    if not isinstance(blob, (bytes, bytearray)) or not blob.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ValueError("보고서 그림이 PNG 형식이 아닙니다.")
    if len(blob) > 8 * 1024 * 1024:
        raise ValueError("보고서 그림 한 장은 8MB를 넘을 수 없습니다.")
    return bytes(blob)


def _v(value: Any, digits: int = 3) -> Any:
    return round(value, digits) if isinstance(value, float) else ("-" if value is None else value)


def _residual(value: Any, ok: Any, tol: float) -> str:
    """평형 검산 잔차 한 줄. 상대오차는 1e-6 급이라 소수 셋째 자리로 반올림하면
    전부 0 이 되어 버린다 — 지수 표기로 쓰고 허용오차·판정을 같이 적는다."""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return "-"
    verdict = "-" if ok is None else ("PASS" if ok else "FAIL")
    return f"{float(value):.2e}  (tolerance {tol:.0e}, {verdict})"


def _first(row: dict, *keys: str) -> Any:
    for key in keys:
        if row.get(key) is not None:
            return row[key]
    return None


def _wrapped_rows(text: str) -> int:
    """병합 폭 안에서 이 글이 실제 몇 줄이 되는지. 한글·기호는 2배 폭으로 센다.

    행 수를 상수로 박아 두면 범위 문구가 한 항목만 늘어도 **조용히 잘린다**
    (실측: ASSESSMENT_SCOPE 10줄이 폭 77자에서 15줄이 되어 10행 상자를 넘쳤다).
    """
    total = 0
    for line in text.split("\n"):
        width = sum(2 if ord(ch) > 0x2000 else 1 for ch in line)
        total += max(1, -(-width // PARA_CHARS))
    return total


def _merge(ws, row: int, c1: int, c2: int, value: Any = None, *,
           fill: str | None = None, bold: bool = False, color: str = SLATE,
           size: float = 9, align: str = "left", border: bool = False,
           wrap: bool = True) -> None:
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = ws.cell(row, c1, _v(value))
    cell.font = Font(name="Arial", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        thin = Side(style="thin", color=LINE)
        for col in range(c1, c2 + 1):
            ws.cell(row, col).border = Border(top=thin, bottom=thin,
                                               left=thin if col == c1 else None,
                                               right=thin if col == c2 else None)


def _chapter(ws, row: int, title: str) -> int:
    _merge(ws, row, 5, 51, title, bold=True, color=NAVY, size=13)
    ws.row_dimensions[row].height = 22
    return row + 3


def _section(ws, row: int, title: str) -> int:
    _merge(ws, row, 5, 51, title, fill=PALE, bold=True, color=NAVY, size=11, border=True)
    ws.row_dimensions[row].height = 19
    return row + 2


def _paragraph(ws, row: int, text: str, rows: int = 2) -> int:
    ws.merge_cells(start_row=row, start_column=6, end_row=row + rows - 1, end_column=50)
    cell = ws.cell(row, 6, text)
    cell.font = Font(name="Arial", size=9, color=SLATE)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return row + rows + 1


def _kv(ws, row: int, label: str, value: Any) -> int:
    _merge(ws, row, 6, 17, label, fill="EEF3F7", bold=True, border=True)
    _merge(ws, row, 18, 50, value, border=True)
    ws.row_dimensions[row].height = 18
    return row + 1


def _table(ws, row: int, spans: list[tuple[int, int]], headers: list[str],
           rows: list[list[Any]], max_rows: int | None = None) -> int:
    for (c1, c2), header in zip(spans, headers):
        _merge(ws, row, c1, c2, header, fill=NAVY, bold=True, color=WHITE,
               align="center", border=True)
    row += 1
    for ridx, values in enumerate(rows if max_rows is None else rows[:max_rows]):
        fill = "F6F8FA" if ridx % 2 else WHITE
        for (c1, c2), value in zip(spans, values):
            _merge(ws, row, c1, c2, value, fill=fill, align="center", border=True)
        ws.row_dimensions[row].height = 17
        row += 1
    return row


def _add_image(ws, png: bytes, anchor_row: int, max_w: int, max_h: int) -> int:
    """F열 `anchor_row` 에 그림을 얹고 **실제로 차지한 행 수**를 돌려준다.

    캡션 행을 상수로 박아 두면 캡처 비율이 바뀔 때마다 캡션이 그림에서 떨어지거나
    겹친다(실측: 16:9 화면 캡처에서 캡션이 그림 아래 20~26행 = 약 400~520px 떠 있었다).
    """
    image = XLImage(io.BytesIO(png))
    scale = min(max_w / image.width, max_h / image.height)
    image.width, image.height = int(image.width * scale), int(image.height * scale)
    ws.add_image(image, f"F{anchor_row}")
    return -(-image.height // DEFAULT_ROW_PX)


def _figure(ws, row: int, image: Optional[bytes], caption: str, *,
            max_w: int, max_h: int, missing_note: str) -> int:
    """그림 한 장 + 캡션. 그림이 없으면 **자리표시 한 줄**로 대체한다.

    예전에는 LC 그림 하나만 없어도 보고서 생성 자체가 실패했다. 렌더가 실패해도
    나머지 결과는 멀쩡하므로 그 칸만 비운다(Unit 권상 보고서와 같은 방침).
    """
    if image is None:
        _merge(ws, row, 6, 50, missing_note, fill="FFF2CC", color=SLATE,
               size=9, align="center", border=True)
        ws.row_dimensions[row].height = 20
        return row + 2
    used = _add_image(ws, image, row, max_w, max_h)
    _merge(ws, row + used + 1, 6, 50, caption, color=SLATE, size=8, align="center")
    return row + used + 2


def _new_page(ws, start: int, title: str) -> int:
    if start > 4:
        ws.row_breaks.append(Break(id=start - 1))
    return _chapter(ws, start, title)


def _configure(ws, subtitle: str, doc_meta: str) -> None:
    ws.title = "Report"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.view = "pageBreakPreview"
    ws.freeze_panes = "E4"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left, ws.page_margins.right = 0.59, 0.20
    ws.page_margins.top, ws.page_margins.bottom = 0.59, 0.59
    ws.page_margins.header, ws.page_margins.footer = 0.39, 0.28
    for col in range(1, LAST_COL + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 1.72
    ws.column_dimensions["A"].width = 1.0
    _merge(ws, 1, 5, 51, "STRENGTH ANALYSIS OF MODULE UNIT FOR OCEAN TRANSPORT",
           bold=True, color=NAVY, size=11)
    _merge(ws, 2, 5, 29, subtitle, color=SLATE, size=8)
    _merge(ws, 2, 30, 51, doc_meta, color=SLATE, size=8, align="right")
    ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height = 17, 12, 6
    ws.oddFooter.center.text = "HiTESS WorkBench · Module Unit Ocean Transport"
    ws.oddFooter.right.text = "Page &[Page] / &[Pages]"


def build_module_ocean_report(
    *, stress_json_path: str, leg_json_path: str | None, weld_json_path: str | None,
    figures: Dict[str, bytes], metadata: dict[str, Any] | None = None,
    figure_warnings: Sequence[str] = (),
) -> tuple[str, bytes, dict[str, Any]]:
    """저장된 결과 JSON + 백엔드가 그린 PNG 로 보고서를 조립한다.

    `figures` 는 `module_ocean_figures.render_report_figures()` 의 반환 그대로다
    (`model_iso.png` / `stress_envelope.png` / `stress_<LC>.png`). 빠진 그림은
    자리표시로 대체하고 경고로 남긴다 — 그림 하나 때문에 보고서를 못 내면 안 된다.
    """
    stress, leg, weld = (_read_json(stress_json_path), _read_json(leg_json_path),
                         _read_json(weld_json_path))
    meta = metadata or {}
    images = {str(name): _png(blob) for name, blob in (figures or {}).items()}
    warnings = [str(item) for item in figure_warnings]

    summary = stress.get("summary") or {}
    displacement = stress.get("displacement") or {}
    disp_summary = displacement.get("summary") or stress.get("displacementSummary") or {}
    quality = stress.get("quality") or {}
    scope = stress.get("assessmentScope") or leg.get("assessmentScope") or {}
    cases = stress.get("loadCases") or []
    stress_lc = stress.get("perLoadCase") or {}
    disp_lc = displacement.get("perLoadCase") or stress.get("displacementPerLoadCase") or {}
    # 6장 뒤에 LC 페이지가 끼므로 7·8장 쪽번호가 LC 수만큼 밀린다. 목차와 본문이
    # 같은 숫자를 쓰도록 **여기서 한 번만** 세고 아래 두 곳이 이것만 참조한다.
    # ⚠ 그림 유무가 아니라 **하중조건 수**로 센다 — 그림 한 장이 렌더에 실패해도
    #   쪽번호가 흔들리면 안 된다(그 칸은 자리표시로 채운다).
    case_figure_pages = [(str(case.get("id")), case) for case in cases
                         if case.get("id")] if len(cases) > 1 else []
    case_page_count = len(case_figure_pages)
    warnings += [f"{cid} 응력 그림을 만들지 못했습니다." for cid, _ in case_figure_pages
                 if f"stress_{cid}.png" not in images]
    weld_summary = weld.get("summary") or {}
    verdict = "FAIL" if int(summary.get("exceedCount") or 0) or weld_summary.get("status") == "NG" else "PASS"
    project = str(meta.get("project") or "Module Unit")
    unit = str(meta.get("moduleUnit") or os.path.splitext(os.path.basename(stress_json_path))[0].replace("_ocean_stress", ""))
    report_date = str(meta.get("reportDate") or date.today().isoformat())

    wb = Workbook()
    ws = wb.active
    _configure(ws, f"{project}   |   {len(cases)} load case(s)   |   {unit}",
               f"RUN {meta.get('runId') or '-'}   |   {report_date}")

    row = _section(ws, 4, "Summary of results")
    result_text = (f"RESULT : {verdict} - maximum usage {_v(summary.get('maxUsage'))}, "
                   f"{int(summary.get('exceedCount') or 0)} member(s) exceed the acceptance criteria")
    _merge(ws, row, 6, 50, result_text, fill=GREEN if verdict == "PASS" else RED,
           bold=True, color=NAVY, size=11, align="center", border=True)
    ws.row_dimensions[row].height = 24
    row += 2
    for label, value in [
        ("Case", project), ("Module Unit", unit),
        ("Structure", f"{summary.get('elementCount', '-')} elements, {stress.get('supportCount', '-')} supports"),
        ("Load cases", f"{len(cases)} subcases, acceleration load including gravity"),
        ("Criteria", f"Allowable stress = {stress.get('allowableMPa', '-')} MPa"),
        ("Solution check", "OK" if quality.get("trustworthy", True) else "REVIEW REQUIRED - solver quality warning"),
        ("Maximum usage", f"{_v(summary.get('maxUsage'))} at EID {summary.get('maxStressElementId', '-')} ({summary.get('governingLoadCase', '-')})"),
        ("Maximum displacement", f"{_v(_first(disp_summary, 'maxMagMm', 'maxMagnitudeMm'))} mm at GRID {disp_summary.get('maxNodeId', '-')}")]:
        row = _kv(ws, row, label, value)
    row += 1
    row = _section(ws, row, "Contents")
    contents = [("1", "Introduction and scope"), ("2", "Analysis method"), ("3", "Analysis model"),
                ("4", "Loads"), ("5", "Solution verification"), ("6", "Results"),
                ("7", "Leg reaction and weld assessment"), ("8", "Conclusion")]
    # 1쪽=요약·목차, 2~7쪽=1~6장. 6장 뒤로 LC 결과 n 쪽 + 포락 부재 표 1쪽이
    # 더 들어가므로 7·8장은 그만큼 뒤로 민다(총 쪽수 10 + n 과 같은 셈법이다).
    chapter_pages = [2, 3, 4, 5, 6, 7, 9 + case_page_count, 10 + case_page_count]
    _table(ws, row, [(6, 10), (11, 45), (46, 50)], ["No.", "Chapter", "Page"],
           [[n, title, page] for (n, title), page in zip(contents, chapter_pages)])

    start = PAGE_ROWS + 3
    row = _new_page(ws, start, "1.  Introduction")
    row = _section(ws, row, "1.1  Objective")
    row = _paragraph(ws, row, "This report documents the linear static strength assessment of the Module Unit and the supporting deck-leg welds under ocean transportation acceleration load cases.", 3)
    row = _section(ws, row, "1.2  Scope")
    scope_lines = ([f"Included - {x}" for x in scope.get("included", [])]
                   + [f"Excluded - {x}" for x in scope.get("excluded", [])])
    scope_text = "\n".join(scope_lines) or "Scope information is not available."
    # 판정 범위는 이 보고서에서 가장 잘리면 안 되는 글이다 — 접히는 줄까지 세어 배정한다.
    row = _paragraph(ws, row, scope_text, max(10, _wrapped_rows(scope_text) + 1))
    row = _section(ws, row, "1.3  Limitation")
    _paragraph(ws, row, str(scope.get("note") or "Items outside the stated scope require separate assessment."), 4)

    start += PAGE_ROWS
    row = _new_page(ws, start, "2.  Analysis method")
    row = _section(ws, row, "2.1  Structural idealisation")
    row = _paragraph(ws, row, "The uploaded Module Unit BDF is combined with the selected built-in transport deck. Support nodes are connected to the nearest deck landing nodes using rigid elements. The original deck-leg SPC set is retained as the global boundary condition.", 5)
    row = _section(ws, row, "2.2  Acceptance criteria")
    row = _kv(ws, row, "Member allowable", f"{stress.get('allowableMPa', '-')} MPa")
    row = _kv(ws, row, "Member criterion", "Axial + bending normal stress; shell von Mises stress")
    row = _kv(ws, row, "Weld criterion", f"Equivalent weld stress / allowable; status {weld_summary.get('status', '-')}")
    row += 2
    row = _section(ws, row, "2.3  Solution and verification")
    _paragraph(ws, row, "MSC Nastran SOL 101 is solved using multiple SUBCASE entries. Member stress, displacement, support reactions and weld demand are enveloped by their own governing load cases.", 4)

    start += PAGE_ROWS
    row = _new_page(ws, start, "3.  Analysis model")
    row = _section(ws, row, "3.1  Model overview")
    _figure(ws, row, images.get("model_iso.png"),
            "Figure 3-1  Module Unit analysis model (isometric view)",
            max_w=620, max_h=385,
            missing_note="Figure 3-1 is not available - the model could not be rendered.")

    start += PAGE_ROWS
    row = _new_page(ws, start, "4.  Loads")
    row = _section(ws, row, "4.1  Ocean transport acceleration load cases")
    case_rows = []
    for case in cases:
        cid, accel = str(case.get("id") or "-"), case.get("accelG") or {}
        s, d = stress_lc.get(cid) or {}, disp_lc.get(cid) or {}
        case_rows.append([cid, case.get("label"), accel.get("ax"), accel.get("ay"), accel.get("az"),
                          s.get("maxStressMPa"), _first(d, "maxMagMm", "maxMagnitudeMm")])
    _table(ws, row, [(6, 10), (11, 18), (19, 24), (25, 30), (31, 36), (37, 43), (44, 50)],
           ["LC", "Signs", "ax (g)", "ay (g)", "az (g)", "Max stress", "Max disp."], case_rows)

    start += PAGE_ROWS
    row = _new_page(ws, start, "5.  Solution verification")
    row = _section(ws, row, "5.1  Solver diagnostics")
    row = _kv(ws, row, "Result quality", "TRUSTWORTHY" if quality.get("trustworthy", True) else "REVIEW REQUIRED")
    row = _kv(ws, row, "High pivot DOF", quality.get("highPivotDofCount", 0))
    row = _kv(ws, row, "Affected elements", quality.get("affectedElementCount", 0))
    row += 2
    row = _section(ws, row, "5.2  Global equilibrium")
    # ⚠ 키 이름은 envelope_leg_reactions() 의 반환과 **글자까지 같아야** 한다.
    #   예전에는 forceResidual/momentResidual/status 를 찾아 세 줄이 모두 '-' 로 비었다.
    #   평형 검산은 이 해석의 신뢰 근거인데 보고서에서 통째로 사라져 있었다.
    #   단일 LC 결과 파일(extract_leg_reactions)은 모멘트를 중첩 dict 로 주므로 둘 다 읽는다.
    check = leg.get("check") or {}
    force_err, force_ok = check.get("maxRelError"), check.get("ok")
    moment_err, moment_ok = check.get("momentMaxRelError"), check.get("momentOk")
    if moment_err is None:
        nested = check.get("moment") or {}
        moment_err, moment_ok = nested.get("maxRelError"), nested.get("ok")
    all_ok = check.get("allOk")
    for label, value in [
        ("Force residual", _residual(force_err, force_ok, REACTION_CHECK_TOL)),
        ("Moment residual", _residual(moment_err, moment_ok, MOMENT_CHECK_TOL)),
        ("Check status", "-" if all_ok is None else (
            "PASS - support reactions balance the applied acceleration load"
            if all_ok else "FAIL - reactions do not balance; see result JSON")),
    ]:
        row = _kv(ws, row, label, value)

    start += PAGE_ROWS
    row = _new_page(ws, start, "6.  Results")
    row = _section(ws, row, "6.1  Governing member stress envelope")
    _figure(ws, row, images.get("stress_envelope.png"),
            f"Figure 6-1  Stress usage envelope - governing {summary.get('governingLoadCase', '-')},"
            f" EID {summary.get('maxStressElementId', '-')}",
            max_w=620, max_h=360,
            missing_note="Figure 6-1 is not available - the stress envelope could not be rendered.")

    # Mooring 보고서처럼 LC마다 결과 그림과 수치 표를 독립 페이지로 반복한다.
    # 포락 한 장만 넣으면 여러 SUBCASE를 풀고도 문서상 LC1만 보이는 문제가 생긴다.
    per_case_top = stress.get("perLoadCaseTopElements") or {}
    for case_index, (cid, case) in enumerate(case_figure_pages, start=1):
        start += PAGE_ROWS
        row = _new_page(ws, start, f"6.{case_index + 1}  {cid} stress result")
        summary_lc = stress_lc.get(cid) or {}
        accel = case.get("accelG") or {}
        row = _kv(ws, row, "Acceleration (g)",
                  f"ax={_v(accel.get('ax'))}, ay={_v(accel.get('ay'))}, az={_v(accel.get('az'))}")
        row = _kv(ws, row, "Maximum stress",
                  f"{_v(summary_lc.get('maxStressMPa'))} MPa at EID {summary_lc.get('maxStressElementId', '-')}")
        row = _kv(ws, row, "Maximum usage", _v(summary_lc.get("maxUsage")))
        next_row = _figure(
            ws, row + 1, images.get(f"stress_{cid}.png"),
            f"Figure 6-{case_index + 1}  {cid} member stress distribution",
            max_w=620, max_h=300,
            missing_note=f"Figure 6-{case_index + 1} ({cid}) is not available - "
                         "this load case could not be rendered.")
        top_lc = per_case_top.get(cid) or []
        if top_lc:
            _table(ws, next_row + 1, [(6, 13), (14, 22), (23, 32), (33, 41), (42, 50)],
                   ["EID", "Element", "Stress (MPa)", "Usage", "Result"],
                   [[x.get("elementId"), x.get("type"), x.get("stressMPa"), x.get("usage"),
                     x.get("verdict") or ("NG" if float(x.get("usage") or 0) > 1 else "OK")]
                    for x in top_lc], 5)

    start += PAGE_ROWS
    row = _new_page(ws, start, f"6.{case_page_count + 2}  Governing members (envelope of all load cases)")
    top = stress.get("topElements") or sorted(stress.get("elements") or [], key=lambda x: float(x.get("usage") or 0), reverse=True)
    top_rows = [[x.get("elementId"), x.get("type"), _first(x, "stressMPa", "maxStressMPa"), x.get("usage"),
                 x.get("loadCase"), x.get("verdict") or ("Fail" if float(x.get("usage") or 0) > 1 else "Pass")]
                for x in top]
    _table(ws, row, [(6, 12), (13, 20), (21, 29), (30, 37), (38, 44), (45, 50)],
           ["EID", "Element", "Stress (MPa)", "Usage", "Governing LC", "Result"], top_rows, 40)

    start += PAGE_ROWS
    row = _new_page(ws, start, "7.  Leg reaction and weld assessment")
    weld_by_leg = {str(_first(x, "legId", "id", "legIndex", "index")): x for x in weld.get("legs") or []}
    leg_rows = []
    for x in leg.get("legs") or []:
        key = str(_first(x, "legId", "id", "legIndex", "index"))
        w = weld_by_leg.get(key) or {}
        leg_rows.append([key, _first(x, "fxN", "fx"), _first(x, "fyN", "fy"), _first(x, "fzN", "fz"),
                         _first(w, "sigmaEqMPa", "equivalentStressMPa"), w.get("usage"), w.get("status")])
    _table(ws, row, [(6, 10), (11, 17), (18, 24), (25, 31), (32, 39), (40, 45), (46, 50)],
           ["Leg", "Fx (N)", "Fy (N)", "Fz (N)", "Weld eq. stress", "Usage", "Result"], leg_rows, 30)

    start += PAGE_ROWS
    row = _new_page(ws, start, "8.  Conclusion")
    conclusion = ("The assessed Module Unit members and deck-leg welds satisfy the stated acceptance criteria."
                  if verdict == "PASS" else
                  "The assessed structure does not satisfy the stated acceptance criteria. Reinforcement or load/support revision is required.")
    _merge(ws, row + 5, 6, 50, conclusion, fill=GREEN if verdict == "PASS" else RED,
           bold=True, color=NAVY, size=11, align="center", border=True)
    ws.row_dimensions[row + 5].height = 34
    _paragraph(ws, row + 10, str(scope.get("note") or "Items outside the assessment scope require separate verification."), 5)

    last_row = start + PAGE_ROWS - 1
    ws.print_area = f"A1:BC{last_row}"
    out = io.BytesIO()
    wb.save(out)
    safe_project = re.sub(r"[^0-9A-Za-z가-힣 _-]+", "", project).strip()
    filename = f"{safe_project or 'ModuleUnit'}_Ocean_Transport_Structural_Report_{date.today().isoformat()}.xlsx"
    pages = 10 + case_page_count
    return filename, out.getvalue(), {
        "verdict": verdict, "sheets": 1, "pages": pages, "figures": len(images),
        # 자리표시로 대체된 그림·렌더 실패를 호출부(라우터 헤더)가 그대로 전달한다.
        "warnings": warnings,
    }
