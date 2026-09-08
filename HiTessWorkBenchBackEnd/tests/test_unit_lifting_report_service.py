import io
import os

import pytest
from openpyxl import load_workbook

from app.services.unit_lifting_report_service import generate_unit_lifting_report

FIX = os.path.join(os.path.dirname(__file__), "fixtures", "unit_lifting_report")
STEM = "3496-35210-A508372_20260108_edit"
INFO = {"nastranResultJson": os.path.join(FIX, f"{STEM}_lifting_nastranResult.json"),
        "stabilityJson": os.path.join(FIX, f"{STEM}_stability.json"),
        "liftingMetaJson": os.path.join(FIX, f"{STEM}_lifting_meta.json"),
        "bdf": os.path.join(FIX, f"{STEM}.bdf")}


def test_generate_returns_named_workbook_and_summary():
    name, data, warnings, summary = generate_unit_lifting_report(
        INFO, {"revision": "A", "department": "구조기본설계부"}, generated_by="A476854")
    assert name.startswith("3496-35210_Hydro_Crane_Module_Unit_권상_구조_검토_보고서_RevA_")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Report", "Members", "Displacements", "Wires"]
    assert summary["structure"] == "OK" and summary["stability"] == "warn" and summary["jigRequired"] is False
    assert summary["figureCount"] >= 10
    assert warnings == []


def test_options_override_identity_and_jig():
    name, _data, _warnings, summary = generate_unit_lifting_report(
        INFO, {"hullNo": "9999", "unitNo": "11111", "jigLimitTon": 0.5}, generated_by="X")
    assert name.startswith("9999-11111_")
    assert summary["jigRequired"] is True


def test_missing_required_json_raises():
    with pytest.raises(FileNotFoundError):
        generate_unit_lifting_report({"nastranResultJson": "nope.json", "stabilityJson": "nope.json"}, {})
