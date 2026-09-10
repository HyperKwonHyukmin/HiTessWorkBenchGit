"""결과 레포트 — 사내 표준 서식(ver2609) 채우기 검증."""
import io
import os

import pytest
from openpyxl import load_workbook

from app.services.unit_lifting_report import result_report as RR
from app.services.unit_lifting_report_service import generate_result_report

FIX = os.path.join(os.path.dirname(__file__), "fixtures", "unit_lifting_report")
STEM = "3496-35210-A508372_20260108_edit"
INFO = {"nastranResultJson": os.path.join(FIX, f"{STEM}_lifting_nastranResult.json"),
        "stabilityJson": os.path.join(FIX, f"{STEM}_stability.json"),
        "liftingMetaJson": os.path.join(FIX, f"{STEM}_lifting_meta.json"),
        "bdf": os.path.join(FIX, f"{STEM}.bdf")}


@pytest.fixture(scope="module")
def built():
    name, data, warnings, summary = generate_result_report(
        INFO, {"department": "구조기본설계부"}, generated_by="A476854")
    return name, load_workbook(io.BytesIO(data)), warnings, summary


def test_file_name_follows_company_convention(built):
    name, _wb, _w, _s = built
    assert name.startswith("3496-35210_Hydro_Crane_Module_Unit_권상_구조_해석_보고서_")
    assert name.endswith(".xlsx")


def test_header_and_summary_cells(built):
    _n, wb, _w, _s = built
    ws = wb["Report"]
    assert ws["O3"].value == "Module Unit 권상 구조 해석 보고서"
    assert ws["O4"].value == "HULL NO. 3496"
    assert ws["Y4"].value == "UNIT NO. 35210"
    assert ws["AI4"].value == "Hydro Crane"
    assert ws["F13"].value == pytest.approx(6.74, abs=1e-3)
    assert ws["N13"].value == pytest.approx(275.0)
    assert ws["AD13"].value == pytest.approx(23.52, abs=1e-2)
    assert ws["AL13"].value == pytest.approx(145.76, abs=1e-2)
    # 허용응력·검토결과는 서식의 수식을 그대로 둔다
    assert str(ws["V13"].value).startswith("=")
    assert str(ws["AT13"].value).startswith("=IF")
    assert "23.52" in ws["N74"].value
    assert "145.8" in ws["N88"].value


def test_hook_table_filled_and_blank_rows_cleared(built):
    _n, wb, _w, _s = built
    ws = wb["Report"]
    assert ws["AD92"].value == pytest.approx(6.2)
    # 그룹·러그는 캡처 그림의 라벨(G1·N34)과 같은 표기여야 표에서 어느 와이어인지 찾을 수 있다
    assert [ws[f"F{r}"].value for r in RR.HOOK_BLOCK_ROWS] == ["G1", "G2", "G3"]
    assert [ws[f"L{r}"].value for r in (92, 93, 94, 95)] == ["G1-N34", "G1-N2873", None, None]
    assert ws["X92"].value == pytest.approx(10198.65 / 9800, abs=1e-3)
    assert ws["R92"].value == pytest.approx(3.547, abs=1e-2)
    # 그룹 1·2 는 wire 가 2개뿐 → 3·4행의 '지그 불요' 수식을 지워야 오해가 없다
    assert ws["AJ94"].value is None and ws["AJ95"].value is None
    assert str(ws["AJ92"].value).startswith("=IF")


def test_support_page_dropped_when_no_support_beams(built):
    _n, wb, _w, summary = built
    ws = wb["Report"]
    assert summary["supportPage"] is False
    # openpyxl 은 행을 지워도 시트 dimension 을 줄이지 않는다. 실제로 값이 사라졌는지와
    # 인쇄 범위·페이지 나누기가 2페이지로 줄었는지를 본다(인쇄 결과를 정하는 것은 이 둘이다).
    leftover = [c.value for row in ws.iter_rows(min_row=RR.SUPPORT_PAGE_FIRST_ROW) for c in row if c.value is not None]
    assert leftover == []
    assert ws.print_area.endswith(f"${RR.SUPPORT_PAGE_FIRST_ROW - 1}")
    assert [b.id for b in ws.row_breaks.brk] == [54]


def test_repeated_print_title_removed(built):
    """서식에 남아 있던 print_title_rows 는 페이지 머리글을 중복 인쇄시킨다."""
    _n, wb, _w, _s = built
    assert wb["Report"].print_title_rows is None


def test_images_are_replaced_with_3d_captures(built):
    _n, wb, _w, summary = built
    ws = wb["Report"]
    assert summary["figureCount"] == 4                 # 가서포트 없음 → 4장
    # 로고 2개(1·2페이지) + 캡처 4장
    assert len(ws._images) == 6
    for im in ws._images:
        assert im._data()[:4] == b"\x89PNG"


def test_summary_values(built):
    _n, _wb, warnings, summary = built
    assert warnings == []
    assert summary["structure"] == "OK"
    assert summary["jigRequired"] is False
    assert summary["maxStressMPa"] == pytest.approx(145.7619)


def test_missing_required_json_raises():
    with pytest.raises(FileNotFoundError):
        generate_result_report({"nastranResultJson": "nope.json", "stabilityJson": "nope.json"}, {})


def test_footer_contact_uses_logged_in_user():
    """바닥글 '문의' 는 보고서를 만든 WorkBench 사용자(이름/직급/부서)다 — 입력 폼에 없는 정보."""
    _n, data, _w, _s = generate_result_report(
        INFO, {"department": "구조기본설계부"}, generated_by="A476854",
        generator={"name": "김도현", "position": "선임연구원", "department": "구조시스템연구실"})
    ws = load_workbook(io.BytesIO(data))["Report"]
    footer = ws["F52"].value
    assert "문의 | 김도현/선임연구원/구조시스템연구실" in footer
    assert "본 보고서는 Hi-TESS WorkBench를 통해 자동 생성되었습니다." in footer


def test_footer_contact_falls_back_to_employee_id(built):
    """사용자 정보가 비어 있으면 사번으로라도 연락처를 남긴다."""
    _n, wb, _w, _s = built
    assert "문의 | A476854" in wb["Report"]["F52"].value
