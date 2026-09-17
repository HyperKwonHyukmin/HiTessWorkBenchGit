"""Module Unit 해상 운송 XLSX 보고서.

★ 픽스처는 **실제 생산 함수의 반환값**으로 만든다.
  예전에는 손으로 적은 {"id","fx","fy","fz"} / "equivalentStressMPa" 를 넣었는데,
  이 철자들은 보고서 `_first()` 의 **폴백**에만 걸린다. 그래서 진짜 철자
  (index/fxN/sigmaEqMPa/maxRelError)는 한 번도 지나가지 않았고,
  5.2 평형 검산 세 줄이 모두 '-' 로 비는 버그가 테스트를 통과해 버렸다.
  아래 헬퍼는 envelope_leg_reactions()·envelope_weld() 를 직접 호출하므로
  생산 코드의 키가 바뀌면 이 테스트가 먼저 깨진다.
"""
import io
import json

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.services.module_ocean_report import build_module_ocean_report
from app.services.module_ocean_results import envelope_leg_reactions, envelope_weld
from app.services.module_ocean_weld import evaluate_weld


def _png() -> bytes:
    """렌더러가 돌려주는 것과 같은 PNG 바이트."""
    buf = io.BytesIO()
    Image.new("RGB", (320, 180), (20, 35, 55)).save(buf, format="PNG")
    return buf.getvalue()


def _leg_rows() -> list:
    """extract_leg_reactions() 가 내는 Leg 행 모양 그대로."""
    return [
        {"index": idx, "jungbanNodeId": 100 + idx, "spcNodeId": 100 + idx,
         "x": 1000.0 * idx, "y": 0.0, "z": 0.0, "zTopMm": 900.0,
         "fxN": 1000.0 * idx, "fyN": 2000.0 * idx, "fzN": 30000.0 * idx,
         "resultantN": 30000.0 * idx,
         "mxNmm": 1.0e6 * idx, "myNmm": 2.0e6 * idx, "mzNmm": 3.0e5 * idx,
         "weldPlaneZMm": -25.0,
         "momentAtSpcNmm": {"mx": 0.0, "my": 0.0, "mz": 0.0}}
        for idx in (1, 2)
    ]


def _leg_case(case_id: str):
    """extract_leg_reactions() 반환 모양. check 의 키 이름이 이 테스트의 핵심이다."""
    return case_id, {
        "schema": "module-ocean-leg/2",
        "totalMassT": 99.93,
        "accelG": {"ax": 0.2, "ay": 0.3, "az": -1.1},
        "legs": _leg_rows(),
        "weldMomentReference": "weld-plane",
        "check": {
            "sumReactionN": [0.0, 0.0, 0.0], "expectedN": [0.0, 0.0, 0.0],
            "maxRelError": 1.23e-6, "ok": True,
            "moment": {"sumMomentNmm": [0.0, 0.0, 0.0], "expectedNmm": [0.0, 0.0, 0.0],
                       "cogMm": [0.0, 0.0, 0.0], "maxRelError": 3.45e-6, "ok": True},
            "allOk": True,
        },
    }


def _real_leg_and_weld(case_ids: list):
    """포락 결과 = 생산 코드가 만든 진짜 dict."""
    weld = envelope_weld([(cid, evaluate_weld(_leg_rows())) for cid in case_ids])
    leg = envelope_leg_reactions([_leg_case(cid) for cid in case_ids], weld)
    return leg, weld


def _stress(case_ids: list) -> dict:
    return {
        "allowableMPa": 220.0,
        "supportCount": 12,
        "summary": {"elementCount": 3500, "maxStressMPa": 138.9, "maxStressElementId": 3509,
                    "maxUsage": 0.631, "exceedCount": 0, "governingLoadCase": case_ids[0]},
        "elements": [{"elementId": 3509, "type": "CBEAM", "stressMPa": 138.9,
                      "usage": 0.631, "loadCase": case_ids[0]}],
        "topElements": [{"elementId": 3509, "type": "CBEAM", "stressMPa": 138.9,
                         "usage": 0.631, "loadCase": case_ids[0], "verdict": "OK"}],
        "loadCases": [{"id": cid, "label": "+X +Y -Z",
                       "accelG": {"ax": 0.2, "ay": 0.3, "az": -1.1}} for cid in case_ids],
        "perLoadCase": {cid: {"maxStressMPa": 138.9, "maxStressElementId": 3509,
                              "maxUsage": 0.631, "loadCase": cid} for cid in case_ids},
        "perLoadCaseTopElements": {
            cid: [{"elementId": 3509, "type": "CBEAM", "stressMPa": 138.9,
                   "usage": 0.631, "verdict": "OK"}] for cid in case_ids},
        "displacement": {"summary": {"maxMagMm": 45.2, "maxNodeId": 2529,
                                     "loadCase": case_ids[0]},
                         "perLoadCase": {cid: {"maxMagMm": 45.2, "maxNodeId": 2529}
                                         for cid in case_ids}},
        "quality": {"trustworthy": True, "highPivotDofCount": 0, "affectedElementCount": 0},
        "assessmentScope": {"included": ["부재 응력"], "excluded": ["좌굴"],
                            "note": "범위 밖 항목은 별도 검토"},
    }


def _write_case(tmp_path, case_ids):
    leg, weld = _real_leg_and_weld(case_ids)
    paths = {}
    for name, payload in (("stress", _stress(case_ids)), ("leg", leg), ("weld", weld)):
        path = tmp_path / f"{name}.json"
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        paths[name] = str(path)
    return paths, leg, weld


def _build(tmp_path, case_ids, *, with_case_figures=True,
           metadata={"project": "3521", "moduleUnit": "Hydro Crane"}):
    paths, _, _ = _write_case(tmp_path, case_ids)
    figures = {"model_iso.png": _png(), "stress_envelope.png": _png()}
    if with_case_figures:
        figures.update({f"stress_{cid}.png": _png() for cid in case_ids})
    return build_module_ocean_report(
        stress_json_path=paths["stress"], leg_json_path=paths["leg"],
        weld_json_path=paths["weld"], figures=figures, metadata=metadata)


def _cells(data: bytes) -> list:
    ws = load_workbook(io.BytesIO(data))["Report"]
    return [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]


# ── 회귀 ① 평형 검산이 실제 키를 읽는다 ──────────────────────────────────

def test_global_equilibrium_rows_read_the_real_check_keys(tmp_path):
    """envelope_leg_reactions 는 maxRelError/momentMaxRelError/allOk 를 준다.

    보고서가 forceResidual/momentResidual/status 를 찾던 시절에는 세 줄이 모두
    '-' 로 비었다. 평형 검산은 이 해석의 신뢰 근거라 비면 안 된다.
    """
    _, data, _ = _build(tmp_path, ["LC1", "LC5"])
    values = [str(v) for v in _cells(data)]

    force = values[values.index("Force residual") + 1]
    moment = values[values.index("Moment residual") + 1]
    status = values[values.index("Check status") + 1]

    assert force.startswith("1.23e-06") and "PASS" in force, force
    assert moment.startswith("3.45e-06") and "PASS" in moment, moment
    assert status.startswith("PASS"), status


def test_global_equilibrium_reports_a_failed_check(tmp_path):
    """검산이 깨지면 보고서도 FAIL 로 적어야 한다 — 빈칸으로 넘어가면 안 된다."""
    paths, leg, _ = _write_case(tmp_path, ["LC1"])
    leg["check"].update({"ok": False, "momentOk": False, "allOk": False})
    (tmp_path / "leg.json").write_text(json.dumps(leg), encoding="utf-8")

    _, data, _ = build_module_ocean_report(
        stress_json_path=paths["stress"], leg_json_path=paths["leg"],
        weld_json_path=paths["weld"],
        figures={"model_iso.png": _png(), "stress_envelope.png": _png()}, metadata={})

    values = [str(v) for v in _cells(data)]
    assert "FAIL" in values[values.index("Force residual") + 1]
    assert values[values.index("Check status") + 1].startswith("FAIL")


# ── 회귀 ② 목차 쪽번호가 LC 페이지를 센다 ────────────────────────────────

@pytest.mark.parametrize("case_ids, with_case_figures, lc_pages", [
    (["LC1"], False, 0),                                    # 단일 조건 — LC 페이지 없음
    (["LC1", "LC5"], True, 2),
    # 그림이 하나도 없어도 쪽수는 같아야 한다 — 렌더 실패가 쪽번호를 흔들면 안 된다.
    (["LC1", "LC5"], False, 2),
    ([f"LC{n}" for n in range(1, 9)], True, 8),             # 실제 포락 8 조건
])
def test_contents_page_numbers_follow_the_real_layout(tmp_path, case_ids,
                                                      with_case_figures, lc_pages):
    """6장 뒤에 LC 페이지 n + 포락 부재 표 1쪽이 끼므로 7·8장이 그만큼 밀린다.

    예전에는 목차가 장 순서(i+2)만 찍어 8 조건에서 9쪽이 어긋났다.
    """
    _, data, meta = _build(tmp_path, case_ids, with_case_figures=with_case_figures)
    values = _cells(data)

    def page_of(title):
        return values[values.index(title) + 1]

    assert page_of("Introduction and scope") == 2
    assert page_of("Results") == 7
    assert page_of("Leg reaction and weld assessment") == 9 + lc_pages
    assert page_of("Conclusion") == 10 + lc_pages
    # 목차 마지막 쪽 == 보고서 총 쪽수. 둘이 어긋나면 어느 한쪽이 거짓말한다.
    assert meta["pages"] == 10 + lc_pages


# ── 기존 커버리지 ────────────────────────────────────────────────────────

def test_build_report_contains_results_and_viewer_captures(tmp_path):
    filename, data, meta = _build(tmp_path, ["LC1", "LC2"])

    assert filename.startswith("3521_Ocean_Transport_Structural_Report_")
    assert meta == {"verdict": "PASS", "sheets": 1, "pages": 12, "figures": 4,
                    "warnings": []}
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Report"]
    ws = wb["Report"]
    assert len(ws._images) == 4
    assert ws.print_title_rows == "$1:$3"
    assert ws.page_setup.orientation == "portrait"
    values = _cells(data)
    for expected in ("Summary of results", "3.  Analysis model",
                     "7.  Leg reaction and weld assessment",
                     "6.2  LC1 stress result", "6.3  LC2 stress result"):
        assert expected in values


def test_leg_table_uses_production_reaction_and_weld_keys(tmp_path):
    """Leg 표가 fxN·sigmaEqMPa 를 실제로 읽는지 — 폴백 철자로 통과하면 안 된다."""
    leg, weld = _real_leg_and_weld(["LC1"])
    expected_fx = leg["legs"][0]["fxN"]
    expected_sigma = round(weld["legs"][0]["sigmaEqMPa"], 3)

    _, data, _ = _build(tmp_path, ["LC1"], with_case_figures=False)
    values = _cells(data)
    assert expected_fx in values, "Leg 반력 Fx 가 표에 없다"
    assert expected_sigma in values, "용접 등가응력이 표에 없다"


def test_missing_figures_become_placeholders_not_a_failure(tmp_path):
    """그림이 빠져도 보고서는 나와야 한다.

    예전에는 LC 그림 하나만 없어도 400 으로 전체가 실패했다. 렌더가 실패해도
    응력·반력·용접 숫자는 멀쩡하므로 그 칸만 자리표시로 비우고 경고로 알린다.
    """
    paths, _, _ = _write_case(tmp_path, ["LC1", "LC5"])
    _, data, meta = build_module_ocean_report(
        stress_json_path=paths["stress"], leg_json_path=paths["leg"],
        weld_json_path=paths["weld"], figures={}, metadata={},
        figure_warnings=["stress_envelope.png: RuntimeError: 렌더 실패"])

    values = [str(v) for v in _cells(data)]
    assert any("Figure 3-1 is not available" in v for v in values)
    assert any("Figure 6-1 is not available" in v for v in values)
    assert meta["figures"] == 0
    # 렌더러가 준 사유 + 빠진 LC 그림이 모두 경고에 남는다.
    assert "stress_envelope.png: RuntimeError: 렌더 실패" in meta["warnings"]
    assert any("LC5" in w for w in meta["warnings"]), meta["warnings"]


def test_a_non_png_figure_is_rejected(tmp_path):
    """렌더러가 PNG 가 아닌 것을 주면 조용히 넣지 말고 막는다."""
    paths, _, _ = _write_case(tmp_path, ["LC1"])
    with pytest.raises(ValueError, match="PNG"):
        build_module_ocean_report(
            stress_json_path=paths["stress"], leg_json_path=None, weld_json_path=None,
            figures={"model_iso.png": b"not a png"}, metadata={})
