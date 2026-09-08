"""ReportSheet — 고정 페이지 틀(머리글·본문·바닥글 + 외곽선) 레이아웃 검증."""
import io

from openpyxl import Workbook, load_workbook

from app.services.unit_lifting_report.sheet import (
    BODY_ROWS, COLS, HEADER_ROWS, PAGE_ROWS, ReportSheet, text_width,
)

PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d494844520000000100000001080200000090775"
    "3de0000000c4944415408d763f8cfc00000030101000ce3b60e0000000049454e44ae426082")


def _sheet(**kw):
    wb = Workbook()
    return wb, ReportSheet(wb.active, doc_title="검토 보고서", hull="3496", unit="35210",
                           method="Hydro Crane", **kw)


def _roundtrip(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf)


def test_text_width_counts_hangul_as_two():
    assert text_width("abc") == 3
    assert text_width("가나다") == 6
    assert text_width("A가") == 3


def test_page_frame_has_header_footer_and_outline():
    wb, s = _sheet(total_pages=2)
    s.para("본문")
    total = s.finish()
    ws = _roundtrip(wb)["Report"]
    assert total == 1
    # 머리글: HULL / UNIT / 권상 방식 이 4번째 행에 들어간다
    assert ws.cell(row=4, column=1).value == "HULL NO. 3496"
    assert ws.cell(row=4, column=5).value == "UNIT NO. 35210"
    assert ws.cell(row=4, column=9).value == "Hydro Crane"
    # 바닥글
    foot = 1 + HEADER_ROWS + BODY_ROWS + 1
    assert ws.cell(row=foot, column=9).value == "Page 1 / 2"
    # 외곽선(medium) 이 프레임 네 변에 있다
    assert ws.cell(row=1, column=1).border.top.style == "medium"
    assert ws.cell(row=PAGE_ROWS, column=1).border.bottom.style == "medium"
    assert ws.cell(row=10, column=1).border.left.style == "medium"
    assert ws.cell(row=10, column=COLS).border.right.style == "medium"


def test_chapter_starts_new_frame_and_records_page():
    wb, s = _sheet()
    s.para("본문")
    p1 = s.chapter("1", "1. 개요")
    s.para("x")
    p2 = s.chapter("2", "2. 모델")
    total = s.finish()
    assert (p1, p2) == (2, 3)
    assert s.chapter_pages == {"1": 2, "2": 3}
    assert total == 3
    assert len(s.ws.row_breaks.brk) == 2          # 페이지 사이에만 나누기


def test_body_overflow_opens_next_page():
    wb, s = _sheet()
    for _ in range(BODY_ROWS + 5):
        s.para("한 줄")
    s.finish()
    assert s.page - 1 == 2                         # 본문이 넘쳐 2페이지가 됐다


def test_table_repeats_header_on_next_page():
    wb, s = _sheet()
    s.table(["A", "B"], [[i, i * 2] for i in range(BODY_ROWS + 10)], widths=[6, 6])
    s.finish()
    ws = _roundtrip(wb)["Report"]
    headers = [c.row for col in ws.iter_cols(min_col=1, max_col=1) for c in col if c.value == "A"]
    assert len(headers) >= 2, "페이지를 넘어가면 표 머리행을 다시 그려야 한다"


def test_table_status_fill_and_number_format():
    wb, s = _sheet()
    s.table(["항목", "값", "판정"], [["a", 1.2345, "OK"], ["b", 2, "NG"]], widths=[4, 4, 4], status_col=2)
    s.finish()
    ws = _roundtrip(wb)["Report"]
    row = next(c.row for col in ws.iter_cols(min_col=1, max_col=1) for c in col if c.value == "a")
    assert ws.cell(row=row, column=5).value == 1.2345
    assert ws.cell(row=row + 1, column=9).value == "NG"
    assert ws.cell(row=row + 1, column=9).fill.fgColor.rgb.endswith("FFE0E3")


def test_toc_two_pass_fills_recorded_rows():
    wb, s = _sheet()
    s.toc_placeholder(2)
    s.chapter("1", "1. 개요")
    s.chapter("2", "2. 모델")
    s.toc_fill([("1. 개요", s.chapter_pages["1"]), ("2. 모델", s.chapter_pages["2"], 1)])
    s.finish()
    ws = _roundtrip(wb)["Report"]
    r0, r1 = s._toc_rows
    assert ws.cell(row=r0, column=1).value == "1. 개요"
    assert ws.cell(row=r0, column=11).value == 2      # 11~12열 병합의 좌상단
    assert ws.cell(row=r1, column=1).value == "2. 모델"
    assert ws.cell(row=r1, column=11).value == 3


def test_figure_shrinks_into_remaining_space():
    wb, s = _sheet()
    for _ in range(BODY_ROWS - 23):               # 잔여 21행: 최소치(20)≤21<필요(23) → 축소 경로
        s.para("한 줄")
    before = s.page
    s.figure(PNG_1PX, "테스트")
    assert s.page == before, "남은 공간이 충분하면 축소해서 같은 페이지에 넣는다"
    assert s.ws._images[-1].height < s.FIG_H
    s.finish()


def test_figure_moves_to_next_page_when_space_too_small():
    wb, s = _sheet()
    for _ in range(BODY_ROWS - 5):
        s.para("한 줄")
    before = s.page
    s.figure(PNG_1PX, "테스트")
    assert s.page == before + 1, "잔여 공간이 최소치보다 좁으면 다음 페이지로 넘긴다"
    s.finish()


def test_cover_occupies_its_own_page():
    wb, s = _sheet()
    s.cover("Module Unit 권상 구조 검토 보고서", [("HULL NO.", "3496")],
            [("구조", "OK", "ok"), ("자세안정성", "WARN", "warn"), ("지그", "불요", "ok")])
    s.para("다음 장 본문")
    s.finish()
    ws = _roundtrip(wb)["Report"]
    vals = [c.value for row in ws.iter_rows(min_row=1, max_row=PAGE_ROWS) for c in row if c.value is not None]
    assert "Module Unit 권상 구조 검토 보고서" in vals
    assert "다음 장 본문" not in vals               # 표지 다음 내용은 새 페이지로
