"""POST /structural-report 라우트.

보고서는 **해석 id 하나**로 만들어진다(Unit 권상 보고서와 같은 규약). 예전에는
결과 파일 경로 3개와 화면 캡처 PNG 를 클라이언트가 올려 보냈고, 그래서 해석 화면이
떠 있어야만 보고서가 나오고 이력에서 다시 뽑을 수 없었다.

이 파일이 지키는 것은 **소유권·상태 검증과 그림 실패의 격리**다. 렌더가 실패해도
숫자는 멀쩡하므로 보고서는 나와야 한다.
"""
import io
import json

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import database, models
from app.dependencies import require_auth
from app.routers import module_ocean_transport as mod
from app.services.module_ocean_results import envelope_leg_reactions, envelope_weld
from app.services.module_ocean_weld import evaluate_weld

OWNER = "OWNER01"


def _client(db_session, employee_id: str) -> TestClient:
    app = FastAPI()
    app.include_router(mod.router)

    def override_db():
        yield db_session

    app.dependency_overrides[database.get_db] = override_db
    app.dependency_overrides[require_auth] = lambda: employee_id
    return TestClient(app)


def _leg_rows():
    return [{"index": i, "jungbanNodeId": 100 + i, "spcNodeId": 100 + i,
             "x": 1000.0 * i, "y": 0.0, "z": 0.0, "zTopMm": 900.0,
             "fxN": 1000.0 * i, "fyN": 2000.0 * i, "fzN": 30000.0 * i,
             "resultantN": 30000.0 * i,
             "mxNmm": 1.0e6 * i, "myNmm": 2.0e6 * i, "mzNmm": 3.0e5 * i,
             "weldPlaneZMm": -25.0,
             "momentAtSpcNmm": {"mx": 0.0, "my": 0.0, "mz": 0.0}} for i in (1, 2)]


def _records(tmp_path, monkeypatch, owner=OWNER):
    """userConnection 아래에 결과 파일 3종을 놓고 result_info 를 만든다."""
    base = tmp_path / "userConnection"
    work_dir = base / f"20260916_120000_{owner}_{mod.PROGRAM_NAME}"
    work_dir.mkdir(parents=True)
    monkeypatch.setattr(mod, "_USER_CONNECTION_DIR", str(base))

    weld = envelope_weld([("LC1", evaluate_weld(_leg_rows()))])
    leg = envelope_leg_reactions([("LC1", {
        "schema": "s", "totalMassT": 99.93,
        "accelG": {"ax": 0.2, "ay": 0.3, "az": -1.1},
        "legs": _leg_rows(), "weldMomentReference": "weld-plane",
        "check": {"maxRelError": 1.23e-6, "ok": True,
                  "moment": {"maxRelError": 3.45e-6, "ok": True}, "allOk": True},
    })], weld)
    stress = {
        "allowableMPa": 220.0,
        "summary": {"elementCount": 1, "maxStressMPa": 138.9, "maxStressElementId": 3509,
                    "maxUsage": 0.631, "exceedCount": 0, "governingLoadCase": "LC1"},
        "elements": [{"elementId": 3509, "type": "CBEAM", "stressMPa": 138.9, "usage": 0.631}],
        "topElements": [{"elementId": 3509, "type": "CBEAM", "stressMPa": 138.9,
                         "usage": 0.631, "verdict": "OK"}],
        "loadCases": [{"id": "LC1", "label": "++-", "accelG": {"ax": .2, "ay": .3, "az": -1.1}}],
        "perLoadCase": {"LC1": {"maxStressMPa": 138.9, "maxUsage": 0.631}},
        "assessmentScope": {"included": ["부재 응력"], "excluded": ["좌굴"], "note": "별도 검토"},
    }
    paths = {}
    for name, payload in (("stress", stress), ("legreact", leg), ("weld", weld)):
        path = work_dir / f"model_{name}.json"
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        paths[name] = str(path)
    return {
        "runId": "job-1",
        # 합본 BDF 는 일부러 두지 않는다 — 그림 없이도 보고서가 나와야 한다.
        "model": {"idOffset": 200000},
        "stress": {"resultJson": paths["stress"], "loadCases": stress["loadCases"]},
        "legReaction": {"resultJson": paths["legreact"]},
        "weld": {"resultJson": paths["weld"]},
    }


def _analysis(db_session, result_info, *, owner=OWNER, status="Success",
              program=None) -> int:
    record = models.Analysis(
        job_id="job-1", project_name="3521", employee_id=owner,
        program_name=program or mod.STRUCTURAL_PROGRAM_NAME,
        status=status, input_info={}, result_info=result_info, source="web",
    )
    db_session.add(record)
    db_session.commit()
    db_session.refresh(record)
    return record.id


def test_report_is_built_from_the_analysis_id_alone(db_session, tmp_path, monkeypatch):
    """클라이언트가 보내는 것은 id 하나뿐이다 — 경로도 캡처도 올리지 않는다."""
    analysis_id = _analysis(db_session, _records(tmp_path, monkeypatch))

    response = _client(db_session, OWNER).post(
        "/api/analysis/module-ocean-transport/structural-report",
        json={"analysis_id": analysis_id},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    ws = load_workbook(io.BytesIO(response.content))["Report"]
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    # 평형 검산이 실제로 실려 있다(옛 키를 읽던 시절에는 '-' 였다).
    assert "1.23e-06" in values[values.index("Force residual") + 1]


def test_missing_combined_bdf_only_warns(db_session, tmp_path, monkeypatch):
    """그림을 못 그려도 보고서는 나온다 — 사유는 요약 헤더에 남긴다."""
    analysis_id = _analysis(db_session, _records(tmp_path, monkeypatch))

    response = _client(db_session, OWNER).post(
        "/api/analysis/module-ocean-transport/structural-report",
        json={"analysis_id": analysis_id},
    )

    assert response.status_code == 200, response.text
    import urllib.parse
    summary = json.loads(urllib.parse.unquote(response.headers["x-report-summary"]))
    assert summary["figures"] == 0
    assert any("합본 BDF" in note for note in summary["warnings"]), summary["warnings"]


def test_another_users_analysis_is_refused(db_session, tmp_path, monkeypatch):
    analysis_id = _analysis(db_session, _records(tmp_path, monkeypatch), owner="SOMEONE")

    response = _client(db_session, OWNER).post(
        "/api/analysis/module-ocean-transport/structural-report",
        json={"analysis_id": analysis_id},
    )
    assert response.status_code == 403


def test_unknown_analysis_is_404(db_session):
    response = _client(db_session, OWNER).post(
        "/api/analysis/module-ocean-transport/structural-report",
        json={"analysis_id": 999999},
    )
    assert response.status_code == 404


def test_other_program_is_refused(db_session, tmp_path, monkeypatch):
    """1단계 검증 레코드로 보고서를 만들면 결과 구조가 달라 엉뚱한 문서가 나온다."""
    analysis_id = _analysis(db_session, _records(tmp_path, monkeypatch),
                            program="GroupModuleUnit")
    response = _client(db_session, OWNER).post(
        "/api/analysis/module-ocean-transport/structural-report",
        json={"analysis_id": analysis_id},
    )
    assert response.status_code == 400


def test_failed_analysis_is_refused(db_session, tmp_path, monkeypatch):
    analysis_id = _analysis(db_session, _records(tmp_path, monkeypatch), status="Failed")
    response = _client(db_session, OWNER).post(
        "/api/analysis/module-ocean-transport/structural-report",
        json={"analysis_id": analysis_id},
    )
    assert response.status_code == 409


def test_missing_stress_result_is_409(db_session, tmp_path, monkeypatch):
    result_info = _records(tmp_path, monkeypatch)
    result_info["stress"] = {}
    analysis_id = _analysis(db_session, result_info)
    response = _client(db_session, OWNER).post(
        "/api/analysis/module-ocean-transport/structural-report",
        json={"analysis_id": analysis_id},
    )
    assert response.status_code == 409
