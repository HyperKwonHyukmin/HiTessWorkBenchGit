"""권상 App 세부 검토의 사용 통계 합산(roll-up) 검증.

ModuleStability · ModuleHoistOptimize · UnitStructuralAnalysis 는 Studio 안의 세부 검토라
통계에서는 부모 App(GroupModuleUnit / SidePassage) 한 줄로 합산되고, 세부 건수는 steps 로 보인다.
세부 검토는 두 App 양쪽에서 생기므로 레코드마다 부모를 찾아야 한다.
"""
from datetime import date, datetime

from app import models
from app.services import usage_rollup
from app.services.usage_report_service import aggregate_period

GMU_DIR = r"C:\HiTessWorkBenchBackEnd\userConnection\20260909_173237_ADMIN001_GroupModuleUnit"
SP_DIR = r"C:\HiTessWorkBenchBackEnd\userConnection\20260915_175754_ADMIN001_SidePassage"


def _add(db, program, *, status="Success", input_info=None, minute=0, employee_id="ADMIN001"):
    rec = models.Analysis(
        employee_id=employee_id,
        program_name=program,
        status=status,
        source="Workbench",
        input_info=input_info or {},
        result_info={},
        created_at=datetime(2026, 9, 20, 9, minute),
    )
    db.add(rec)
    db.commit()
    return rec


def _seed_lifting(db):
    """GMU 부모 1 + 세부 4(자세 2 · 최적화 1 · 구조 1), SidePassage 부모 1 + 세부 2."""
    gmu = _add(db, "GroupModuleUnit", input_info={"bdf_model": GMU_DIR + r"\m.bdf"}, minute=0)
    sp = _add(db, "SidePassage", input_info={"bdf_model": SP_DIR + r"\s.bdf"}, minute=1)
    _add(db, "ModuleStability", input_info={"posture": GMU_DIR + r"\m_posture.json"}, minute=2)
    _add(db, "ModuleStability", status="Failed", input_info={"posture": GMU_DIR + r"\m_posture.json"}, minute=3)
    _add(db, "ModuleHoistOptimize", input_info={"posture": GMU_DIR + r"\m_posture.json"}, minute=4)
    _add(db, "UnitStructuralAnalysis", input_info={"parent_analysis_id": gmu.id}, minute=5)
    _add(db, "ModuleStability", input_info={"posture": SP_DIR + r"\s_posture.json"}, minute=6)
    _add(db, "UnitStructuralAnalysis", input_info={"parent_analysis_id": sp.id}, minute=7)
    _add(db, "BDF Scanner", minute=8)
    return gmu, sp


# ── 판별 규칙 ──────────────────────────────────────────────────────

def test_rollup_resolves_parent_from_posture_folder_and_parent_id(db_session):
    gmu, sp = _seed_lifting(db_session)
    rows = db_session.query(models.Analysis).all()
    parents = usage_rollup.load_parent_programs(db_session, rows)
    by_min = {r.created_at.minute: usage_rollup.rollup_program(r, parents) for r in rows}

    assert by_min[2] == "GroupModuleUnit"  # posture 폴더
    assert by_min[5] == "GroupModuleUnit"  # parent_analysis_id
    assert by_min[6] == "SidePassage"
    assert by_min[7] == "SidePassage"
    assert by_min[8] == "BDF Scanner"      # 그 밖의 App 은 그대로


def test_unresolvable_substep_keeps_own_name():
    rec = models.Analysis(program_name="ModuleStability", input_info={"posture": r"C:\tmp\x_posture.json"})
    assert usage_rollup.rollup_program(rec, {}) == "ModuleStability"
    orphan = models.Analysis(program_name="UnitStructuralAnalysis", input_info={"parent_analysis_id": 999})
    assert usage_rollup.rollup_program(orphan, {}) == "UnitStructuralAnalysis"


def test_posture_folder_with_forward_slashes():
    rec = models.Analysis(
        program_name="ModuleHoistOptimize",
        input_info={"posture": "/srv/userConnection/20260702_151550_A476854_GroupModuleUnit/a_posture.json"},
    )
    assert usage_rollup.rollup_program(rec, {}) == "GroupModuleUnit"


# ── Analysis Management 요약 ─────────────────────────────────────

def test_management_summary_rolls_substeps_into_parent(admin_client, db_session):
    _seed_lifting(db_session)
    summary = admin_client.get("/api/analysis/all", params={"include_summary": True, "limit": 1}).json()["summary"]
    rows = {r["name"]: r for r in summary["programRows"]}

    assert set(rows) == {"GroupModuleUnit", "SidePassage", "BDF Scanner"}
    assert rows["GroupModuleUnit"]["count"] == 5
    assert rows["SidePassage"]["count"] == 3
    assert summary["activePrograms"] == 3
    assert summary["total"] == 9, "레코드 한 건 한 건은 모두 총계에 들어간다"

    steps = {s["label"]: s for s in rows["GroupModuleUnit"]["steps"]}
    assert [s["label"] for s in rows["GroupModuleUnit"]["steps"]] == [
        "모델 입력·검증", "자세안정성 평가", "권상 위치 최적화", "Unit 구조 해석",
    ]
    assert steps["자세안정성 평가"]["count"] == 2
    assert steps["자세안정성 평가"]["success"] == 1
    assert rows["BDF Scanner"]["steps"] == []


# ── 프로그램 상세 모달 ─────────────────────────────────────────────

def test_program_detail_total_matches_summary_row(admin_client, db_session):
    _seed_lifting(db_session)
    body = admin_client.get("/api/analysis/stats/program/GroupModuleUnit").json()

    assert body["summary"]["total"] == 5
    assert body["summary"]["success"] == 4
    assert {s["label"]: s["count"] for s in body["stepBreakdown"]} == {
        "모델 입력·검증": 1, "자세안정성 평가": 2, "권상 위치 최적화": 1, "Unit 구조 해석": 1,
    }
    steps_in_records = sorted(r["step"] for r in body["records"])
    assert steps_in_records.count("자세안정성 평가") == 2


def test_program_detail_for_plain_app_has_no_steps(admin_client, db_session):
    _seed_lifting(db_session)
    body = admin_client.get("/api/analysis/stats/program/BDF%20Scanner").json()
    assert body["summary"]["total"] == 1
    assert body["stepBreakdown"] == []
    assert body["records"][0]["step"] is None


# ── 대시보드 Top 프로그램 ─────────────────────────────────────────

def test_top_programs_counts_substeps_under_parent(admin_client, db_session):
    _seed_lifting(db_session)
    rows = admin_client.get("/api/analysis/stats/top-programs", params={"days": 0}).json()
    counts = {r["program_name"]: r["count"] for r in rows}

    assert counts == {"GroupModuleUnit": 5, "SidePassage": 3, "BDF Scanner": 1}
    assert rows[0]["program_name"] == "GroupModuleUnit"


# ── Usage Reports ────────────────────────────────────────────────

def test_usage_report_rolls_up_and_lists_steps(db_session):
    _seed_lifting(db_session)
    result = aggregate_period(
        db_session, "daily",
        start=datetime(2026, 9, 20, 0, 0, 0),
        end=datetime(2026, 9, 20, 23, 59, 59, 999999),
    )
    programs = {p["name"]: p for p in result["programs"]}

    assert result["total"] == 9
    assert result["activePrograms"] == 3
    assert result["busiestProgram"] == "GroupModuleUnit"
    assert programs["GroupModuleUnit"]["count"] == 5
    assert {s["label"]: s["count"] for s in programs["SidePassage"]["steps"]} == {
        "모델 입력·검증": 1, "자세안정성 평가": 1, "Unit 구조 해석": 1,
    }


def test_usage_report_xlsx_has_steps_and_rollup_columns(db_session):
    from openpyxl import load_workbook
    from app.services.usage_report_service import build_report_xlsx, resolve_period, compute_deltas

    _seed_lifting(db_session)
    bounds = resolve_period("daily", date(2026, 9, 20), today=date(2026, 9, 30))
    cur = aggregate_period(db_session, "daily", bounds.start, bounds.end)
    prev = aggregate_period(db_session, "daily", bounds.prev_start, bounds.prev_end)
    wb = load_workbook(build_report_xlsx(bounds, cur, prev, compute_deltas(cur, prev)))

    programs = list(wb["Programs"].iter_rows(values_only=True))
    assert programs[0][-1] == "세부 검토"
    gmu_row = next(r for r in programs if r[1] == "GroupModuleUnit")
    assert "자세안정성 평가 2" in gmu_row[-1]

    raw = list(wb["Raw Data"].iter_rows(values_only=True))
    assert raw[0][2:4] == ("프로그램", "통계 프로그램")
    assert ("ModuleStability", "SidePassage") in {(r[2], r[3]) for r in raw[1:]}


def test_usage_report_route_keeps_steps_through_response_model(admin_client, db_session):
    """response_model(UsageReportResponse) 이 steps 를 떨구면 화면에 세부 검토가 안 보인다."""
    _seed_lifting(db_session)
    body = admin_client.get("/api/analysis/report", params={"period": "daily", "date": "2026-09-20"}).json()
    programs = {p["name"]: p for p in body["programs"]}

    assert "ModuleStability" not in programs
    assert "ModuleHoistOptimize" not in programs
    assert "UnitStructuralAnalysis" not in programs
    assert programs["GroupModuleUnit"]["count"] == 5
    assert {s["label"]: s["count"] for s in programs["GroupModuleUnit"]["steps"]}["자세안정성 평가"] == 2
