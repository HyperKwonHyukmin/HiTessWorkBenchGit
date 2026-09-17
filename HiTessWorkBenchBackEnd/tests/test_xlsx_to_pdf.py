"""xlsx → PDF 변환기 (Windows + MS Excel COM).

Excel 이 없는 환경(CI·리눅스)에서는 변환 테스트를 건너뛰고, 그 경우에도 사용자에게
사유가 전달되는지(PdfConversionError)만 검증한다.
"""

import io

import pytest
from openpyxl import Workbook

from app.services.xlsx_to_pdf import PdfConversionError, convert_xlsx_to_pdf


def _excel_available() -> bool:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False
    try:
        pythoncom.CoInitialize()
        xl = win32com.client.DispatchEx("Excel.Application")
        xl.Quit()
        return True
    except Exception:
        return False
    finally:
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


needs_excel = pytest.mark.skipif(not _excel_available(), reason="MS Excel 미설치 환경")


def _workbook_bytes(extra_sheet: bool = False) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "권상 구조 검토"
    if extra_sheet:
        data = wb.create_sheet("Members")
        data["A1"] = "데이터 시트 — 인쇄 대상이 아니다"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_empty_input_rejected():
    with pytest.raises(PdfConversionError):
        convert_xlsx_to_pdf(b"")


@needs_excel
def test_converts_to_pdf_bytes():
    data = convert_xlsx_to_pdf(_workbook_bytes())
    # ★ DRM 이 디스크 파일을 감싸므로, 변환기가 Excel 프로세스 안에서 평문으로 되읽어야 한다.
    assert data.startswith(b"%PDF"), data[:16]


@needs_excel
def test_exports_named_sheet_only():
    """워크북째 내보내면 데이터 시트까지 붙어 페이지가 늘어난다 — Report 시트만 나와야 한다."""
    fitz = pytest.importorskip("fitz")
    both = convert_xlsx_to_pdf(_workbook_bytes(extra_sheet=True), sheet_name="Report")
    assert fitz.open(stream=both, filetype="pdf").page_count == 1


@needs_excel
def test_broken_workbook_raises():
    with pytest.raises(PdfConversionError):
        convert_xlsx_to_pdf(b"not a workbook at all")
