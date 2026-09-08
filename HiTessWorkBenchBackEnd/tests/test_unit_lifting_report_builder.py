import io
import os
import shutil

import pytest
from openpyxl import load_workbook

from app.services.unit_lifting_report.collector import collect, ReportOptions
from app.services.unit_lifting_report import figures
from app.services.unit_lifting_report.builder import build_workbook, toc_entries

FIX = os.path.join(os.path.dirname(__file__), "fixtures", "unit_lifting_report")
STEM = "3496-35210-A508372_20260108_edit"


def _info(folder=FIX):
    return {"nastranResultJson": os.path.join(folder, f"{STEM}_lifting_nastranResult.json"),
            "stabilityJson": os.path.join(folder, f"{STEM}_stability.json"),
            "liftingMetaJson": os.path.join(folder, f"{STEM}_lifting_meta.json"),
            "bdf": os.path.join(folder, f"{STEM}.bdf")}


@pytest.fixture(scope="module")
def built():
    data = collect(_info(), ReportOptions(author="A476854", department="구조기본설계부", drawing_no="D-001"), "2026-09-07 15:00")
    figs, errors = figures.render_all_safe(data)
    wb = build_workbook(data, figs, errors)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return data, load_workbook(buf)


def test_sheets_present(built):
    _, wb = built
    assert wb.sheetnames == ["Report", "Members", "Displacements", "Wires"]
    assert wb["Members"].max_row == 61            # 헤더 + 60
    assert wb["Wires"].max_row == 8


def test_cover_summary_and_verdict_text(built):
    _data, wb = built
    joined = "\n".join(str(c.value) for row in wb["Report"].iter_rows() for c in row if c.value is not None)
    assert "Module Unit 권상 구조 검토 보고서" in joined
    assert "HULL NO." in joined and "3496" in joined
    assert "허용응력을 초과하는 부재는 없습니다" in joined
    assert "145.76" in joined and "23.52" in joined
    # fixture 는 JSON 7종이 모두 있으므로 본문 절은 '자료 없음' 이 아니어야 한다.
    # (F06 원본은 2.5MB 라 fixture 에 넣지 않아 부록 C 만 '자료 없음' 이 정상)
    assert "편집 모델(_edited.json)" not in joined
    assert "권상 위치 최적화" not in joined
    assert "6단계 전도 평가 결과" not in joined
    assert joined.count("자료 없음") == 1 and "F06 파일이 결과 폴더에 없어" in joined


def test_toc_pages_monotonic(built):
    """목차는 (제목, 쪽번호) 쌍으로 채워지고 장 쪽번호는 증가해야 한다."""
    data, wb = built
    ws = wb["Report"]
    entries = toc_entries(data)
    # '목차 (Contents)' 제목 다음의 연속 행이 목차다(본문 표에도 11열 정수가 있어 전체 스캔은 못 쓴다).
    head = next(r for r in range(1, ws.max_row + 1)
                if str(ws.cell(row=r, column=1).value or "").strip() == "목차 (Contents)")
    start = next(r for r in range(head + 1, head + 6) if isinstance(ws.cell(row=r, column=11).value, int))
    toc = [(str(ws.cell(row=start + k, column=1).value or "").strip(), ws.cell(row=start + k, column=11).value)
           for k in range(len(entries))]
    assert all(isinstance(p, int) for _t, p in toc), toc
    for (text, level), (toc_text, page) in zip(entries, toc):
        assert toc_text == text, f"목차 제목 불일치: {toc_text!r} != {text!r}"
        assert page > 0, f"{text} 쪽번호가 채워지지 않았다"
    chapter_pages = [p for (_t, lvl), (_tt, p) in zip(entries, toc) if lvl == 0]
    assert chapter_pages == sorted(chapter_pages)


def test_missing_optional_json_writes_placeholder(tmp_path):
    for name in ("_lifting_nastranResult.json", "_stability.json", "_lifting_meta.json"):
        shutil.copy(os.path.join(FIX, f"{STEM}{name}"), tmp_path / f"{STEM}{name}")
    data = collect(_info(str(tmp_path)), ReportOptions())
    figs, errors = figures.render_all_safe(data)
    wb = build_workbook(data, figs, errors)
    joined = "\n".join(str(c.value) for row in wb["Report"].iter_rows() for c in row if c.value is not None)
    assert "자료 없음" in joined
    assert wb["Members"].max_row == 61
