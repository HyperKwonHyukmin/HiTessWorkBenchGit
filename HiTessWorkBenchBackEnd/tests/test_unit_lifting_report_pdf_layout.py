"""결과 레포트 인쇄(PDF) 레이아웃 — 그림 늘림 방지·서식 글자 맞춤·가운데 정렬, 가서포트 그림 배관 반투명.

배경(2026-09-17): 서식의 캡처 자리가 TwoCellAnchor 라 Excel 인쇄에서 그림이 옆으로 11% 늘어났고
(좁은 열의 인쇄 폭 계산 차이), 주의사항 셋째 줄이 박스 밖으로 넘쳐 잘렸다. 여기서는 워크북 안의
앵커·정렬·인쇄 설정을 본다 — 실제 PDF 픽셀 검증은 Excel 이 있는 PC 에서 수동으로 했다.
"""
import io
import json
import math
import os
import shutil

import numpy as np
import pytest
from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor

from app.services.unit_lifting_report import collector as C
from app.services.unit_lifting_report import figures3d, result_report as RR, solid3d as S3
from app.services.unit_lifting_report_service import generate_result_report

FIX = os.path.join(os.path.dirname(__file__), "fixtures", "unit_lifting_report")
STEM = "3496-35210-A508372_20260108_edit"


def _info(folder):
    return {"nastranResultJson": os.path.join(folder, f"{STEM}_lifting_nastranResult.json"),
            "stabilityJson": os.path.join(folder, f"{STEM}_stability.json"),
            "liftingMetaJson": os.path.join(folder, f"{STEM}_lifting_meta.json"),
            "editedJson": os.path.join(folder, f"{STEM}_edited.json"),
            "originalJson": os.path.join(folder, f"{STEM}.json"),
            "bdf": os.path.join(folder, f"{STEM}.bdf")}


@pytest.fixture(scope="module")
def with_supports(tmp_path_factory):
    """fixture 에는 가서포트가 없다 — 편집 모델에 L100×100×10 가서포트 2개를 넣은 사본을 만든다."""
    folder = tmp_path_factory.mktemp("supfix")
    for f in os.listdir(FIX):
        if f.startswith(STEM):
            shutil.copy(os.path.join(FIX, f), folder)
    p = folder / f"{STEM}_edited.json"
    j = json.loads(p.read_text(encoding="utf-8"))
    nodes = {int(n["id"]): n for n in j["nodes"]}
    els = j["elements"]
    stru = sorted({e[k] for e in els if e.get("category") == "Structure" for k in ("startNode", "endNode")})
    connected = {(min(e["startNode"], e["endNode"]), max(e["startNode"], e["endNode"])) for e in els}

    def dist(a, b):
        A, B = nodes[a], nodes[b]
        return math.dist((A["x"], A["y"], A["z"]), (B["x"], B["y"], B["z"]))

    pairs = []
    for a in stru[::7]:
        for b in stru[::11]:
            if a < b and (a, b) not in connected and 900 < dist(a, b) < 1400:
                pairs.append((a, b))
                break
        if len(pairs) == 2:
            break
    assert len(pairs) == 2
    pid = max(x["id"] for x in j["properties"]) + 1
    j["properties"].append({"id": pid, "card": "PBEAML", "kind": "L", "dims": [100, 100, 10, 10],
                            "materialId": j["materials"][0]["id"]})
    eid = max(e["id"] for e in els) + max((r["id"] for r in j.get("rigids", [])), default=0)
    for k, (a, b) in enumerate(pairs):
        els.append({"id": eid + 1 + k, "type": "CBEAM", "startNode": a, "endNode": b, "propertyId": pid,
                    "orientation": [0, 0, 1], "category": "Structure", "modelPart": "stru",
                    "remark": C.SUPPORT_REMARK})
    p.write_text(json.dumps(j, ensure_ascii=False), encoding="utf-8")
    return str(folder)


@pytest.fixture(scope="module")
def built(with_supports):
    _name, data, warnings, summary = generate_result_report(_info(with_supports), {}, generated_by="A476854")
    return load_workbook(io.BytesIO(data)), warnings, summary


# ── 그림 앵커 ────────────────────────────────────────────────────────────────
def test_all_images_use_absolute_anchor(built):
    """TwoCellAnchor 가 하나도 남지 않아야 인쇄에서 그림이 셀 박스에 늘어나지 않는다."""
    wb, _w, summary = built
    ws = wb["Report"]
    assert summary["supportPage"] is True
    assert len(ws._images) == 3 + 6                       # 로고 3 + 캡처 6
    assert not any(isinstance(im.anchor, TwoCellAnchor) for im in ws._images)
    assert all(isinstance(im.anchor, OneCellAnchor) and im.anchor.ext is not None for im in ws._images)


def test_figure_keeps_source_aspect_and_fits_slot(built):
    """앵커 크기(ext)는 원본 PNG 비율 그대로이고, 서식 자리(셀 박스) 안에 들어간다."""
    wb, _w, _s = built
    ws = wb["Report"]
    tpl = RR._load_template()["Report"]
    slots = list(tpl._images)[RR.LOGO_COUNT:]              # 서식의 캡처 자리(TwoCellAnchor)
    figs = list(ws._images)[3:]
    for fig, slot in zip(figs, slots):
        w_px = fig.anchor.ext.width / RR._EMU_PER_PX
        h_px = fig.anchor.ext.height / RR._EMU_PER_PX
        assert w_px / h_px == pytest.approx(fig.width / fig.height, rel=0.01)
        box_w, box_h = RR._cell_box_px(tpl, slot.anchor)
        assert w_px <= box_w + 1 and h_px <= box_h + 1
        assert fig.anchor._from.row == slot.anchor._from.row  # 같은 자리
        # 남는 폭은 양쪽으로 나눠 가운데 — 시작 열 오프셋 = 자리 오프셋 + (박스-그림)/2
        expected = slot.anchor._from.colOff + ((box_w - int(w_px)) // 2) * RR._EMU_PER_PX
        assert fig.anchor._from.colOff == pytest.approx(expected, abs=RR._EMU_PER_PX)


def test_cell_box_px_matches_excel_measurement():
    """실측: 서식의 1.71자 열 = 12px(9.0pt), 13.5pt 행 = 18px. 1페이지 캡처 자리 F16:AZ33 = 46열 × 17행."""
    tpl = RR._load_template()["Report"]
    slot = list(tpl._images)[RR.LOGO_COUNT]                # modelTop
    box_w, box_h = RR._cell_box_px(tpl, slot.anchor)
    assert RR._col_px(tpl, 6) == 12
    assert RR._row_px(tpl, 16) == 18
    # 앵커의 시작·끝 오프셋(EMU) 만큼 한 열·한 행 안쪽으로 들어올 수 있다
    assert abs(box_w - 46 * 12) < 12
    assert abs(box_h - 17 * 18) < 18


# ── 서식 글자 맞춤 · 인쇄 설정 ──────────────────────────────────────────────
def test_form_text_shrinks_to_fit_instead_of_overflowing(built):
    wb, _w, _s = built
    ws = wb["Report"]
    for coord in ("O3", "AS3", "O4", "Y4", "AI4", "F7", "F8", "F9"):
        assert ws[coord].alignment.shrink_to_fit is True, coord
        assert not ws[coord].alignment.wrap_text, coord
    assert ws["F13"].alignment.shrink_to_fit is not True   # 표 숫자칸은 손대지 않는다


def test_print_is_centered_with_visible_content_margins(built):
    wb, _w, _s = built
    ws = wb["Report"]
    assert ws.print_options.horizontalCentered is True
    assert (ws.page_margins.left, ws.page_margins.right) == (0.22, 0.56)
    assert ws.page_setup.scale == 97                        # 100% 면 1페이지가 2pt 넘쳐 쪼개진다


# ── 상세 레포트(sheet.py): 그림은 원본 비율, 주어진 크기는 최대 상자 ─────────────
def test_detail_sheet_figure_keeps_png_aspect():
    """예전엔 640×410 을 표시 크기로 그대로 써서 3D(1.73)·2D(1.44) 그림이 ±10% 눌렸다."""
    from openpyxl import Workbook
    from PIL import Image
    from app.services.unit_lifting_report import sheet as SH

    def png(w, h):
        buf = io.BytesIO()
        Image.new("RGB", (w, h), "white").save(buf, format="PNG")
        return buf.getvalue()

    wb = Workbook()
    s = SH.ReportSheet(wb.active, doc_title="t", hull="1", unit="2", method="Hydro Crane",
                       department="", drawing_no="", contact="", logo_png=None, total_pages=1)
    s.figure(png(1428, 825), "3D")          # 1.731 — 폭 640 에 맞추면 높이 369
    s.figure(png(1080, 750), "2D")          # 1.440 — 높이 410 에 걸려 폭 590
    s.figure(png(1080, 480), "wide", 640, 285)
    got = [(im.width, im.height) for im in s.ws._images]
    for (w, h), src in zip(got, ((1428, 825), (1080, 750), (1080, 480))):
        assert w / h == pytest.approx(src[0] / src[1], rel=0.01), (w, h, src)
    assert got[0] == (640, 369) and got[1] == (590, 410) and got[2] == (640, 284)


# ── 가서포트 그림: 배관 반투명 ───────────────────────────────────────────────
def test_face_colors_returns_rgba_only_when_alpha_given():
    quads = np.zeros((3, 4, 3))
    mesh = S3.SolidMesh(quads, np.ones(3), np.array([1, 2, 3]))
    rgb = S3.face_colors(mesh, per_element={1: (1, 0, 0)})
    assert rgb.shape == (3, 3)
    rgba = S3.face_colors(mesh, per_element={1: (1, 0, 0)}, alpha_per_element={2: 0.2})
    assert rgba.shape == (3, 4)
    assert list(rgba[:, 3]) == [1.0, 0.2, 1.0]


def test_support_figure_dims_pipes_but_not_supports(with_supports, monkeypatch):
    data = C.collect(_info(with_supports), C.ReportOptions.from_payload({}))
    assert len(data.edits.support_beams) == 2
    seen = {}
    real = S3.render

    def spy(mesh, colors, view="home", **kw):
        seen["colors"] = np.asarray(colors)
        seen["elem_ids"] = np.asarray(mesh.elem_ids)
        seen["legend"] = [h.get_label() for h in kw.get("legend", ())]
        return real(mesh, colors, view, **kw)

    monkeypatch.setattr(figures3d.S3, "render", spy)
    png = figures3d.render_support(data, "home")
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    colors, ids = seen["colors"], seen["elem_ids"]
    assert colors.shape[1] == 4
    pipe_ids = set(figures3d._pipe_colors(data))
    assert pipe_ids, "fixture 에 배관이 있어야 의미 있는 검증이다"
    is_pipe = np.isin(ids, list(pipe_ids))
    assert np.allclose(colors[is_pipe, 3], figures3d.PIPE_ALPHA_SUPPORT)
    assert np.allclose(colors[~is_pipe, 3], 1.0)             # 구조·가서포트·권상 장구는 불투명
    assert "배관 (반투명)" in seen["legend"]


def test_other_figures_stay_opaque(with_supports, monkeypatch):
    """반투명은 가서포트 배치 그림에만 — 권상 배치 그림은 종전대로 RGB."""
    data = C.collect(_info(with_supports), C.ReportOptions.from_payload({}))
    shapes = []
    real = S3.render

    def spy(mesh, colors, view="home", **kw):
        shapes.append(np.asarray(colors).shape[1])
        return real(mesh, colors, view, **kw)

    monkeypatch.setattr(figures3d.S3, "render", spy)
    figures3d.render_hoist_model(data, "top")
    assert shapes == [3]


# ── 페이지 외곽선이 모든 페이지에서 한 줄로 이어지는가 ───────────────────────────
def test_result_frame_has_no_hairline(built):
    """서식의 좌우 외곽선은 hairline(0.12pt) 이라 PDF 에서 배율에 따라 사라졌다 나타났다 한다.

    머리글 블록(3~4행)은 thin 이라 한 선의 굵기가 페이지 중간에서 8배 바뀌기도 했다.
    같은 선이 열 스타일과 셀 테두리 두 군데에 나뉘어 있어 둘 다 올려야 한다.
    """
    wb, _w, _s = built
    ws = wb["Report"]
    for letter, edge in (("D", "right"), ("BC", "left")):
        side = getattr(ws.column_dimensions[letter].border, edge)
        assert side.style == "thin", (letter, edge, side.style)
    for col, edge in RR._FRAME_EDGES:
        for row in range(1, ws.max_row + 1):
            side = getattr(ws.cell(row, col).border, edge)
            assert side.style != "hair", f"{ws.cell(row, col).coordinate} {edge} 가 아직 hairline"


def test_detail_images_clear_the_page_outline():
    """상세 레포트의 그림·로고를 A열 왼쪽 끝에 붙이면 흰 배경이 외곽선(medium)을 덮어

    **그림이 있는 페이지에서만** 세로 외곽선이 토막난다(실측: 21쪽 중 7쪽). 앵커를 안으로
    밀어 피한다 — 외곽선은 눈금선을 걸치고 그려져 셀 안쪽으로 1.44pt(≈2px) 들어온다.
    """
    import io as _io
    from PIL import Image
    from openpyxl import Workbook
    from app.services.unit_lifting_report import sheet as SH

    def png(w, h):
        buf = _io.BytesIO()
        Image.new("RGB", (w, h), "white").save(buf, format="PNG")
        return buf.getvalue()

    wb = Workbook()
    s = SH.ReportSheet(wb.active, doc_title="t", hull="1", unit="2", method="Hydro Crane",
                       department="", drawing_no="", contact="", logo_png=png(218, 36), total_pages=1)
    s.figure(png(1428, 825), "3D")
    s.finish()
    assert len(s.ws._images) >= 2                      # 로고 + 그림
    for im in s.ws._images:
        anchor = im.anchor
        assert isinstance(anchor, OneCellAnchor), type(anchor)
        assert anchor._from.col == 0                   # 자리는 그대로 A열
        assert anchor._from.colOff >= 2 * SH.EMU_PER_PX, anchor._from.colOff
        assert anchor.ext.width == int(im.width) * SH.EMU_PER_PX
        assert anchor.ext.height == int(im.height) * SH.EMU_PER_PX
